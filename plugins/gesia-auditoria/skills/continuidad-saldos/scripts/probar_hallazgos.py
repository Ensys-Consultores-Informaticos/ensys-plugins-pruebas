"""Comprueba que el comparador detecta lo que tiene que detectar, y que el papel
lleva el bloque de uso de IA como debe.

    python probar_hallazgos.py --diario "<real.smn>" --cierre cierre.json

Parte de un diario REAL, hace una copia y le mete tres casos a proposito:

  1. una cuenta de resultados en el asiento de apertura (710000000, 333.333)
  2. una cuenta de balance nueva que el expediente no reconoce (122000000, 50.000)
  3. las contrapartidas de las dos, que producen diferencias en cuentas que si
     estan en el cierre (570 y 572)

Por que hace falta: en un expediente sin diferencias, un comparador roto y uno
correcto dan el mismo resultado. Que un expediente salga a cero no demuestra que
esto funcione, solo que no se inventa hallazgos.

Desde la 1.25.0 del MCP la prueba no lee el diario: lee lo que exporta el MCP, que es el
filtro de apertura y el asiento de apertura. El arnes hace esas dos exportaciones con las
MISMAS consultas (SQL_FILTRO y sql_apertura) por ODBC, como haria el MCP, y ademas
comprueba sobre el diario real entero que la condicion traducida a JET selecciona
exactamente lo mismo que el filtro evaluado en pandas, que es como se hacia antes.

Sigue exigiendo un .smn real, igual que el resto de las pruebas del proyecto.
Un fixture sintetico desde cero esta pendiente (fase 3). El primer trozo ya esta: las
diferencias que se compensan (tanda del 21/10/2026) se prueban con filtro, apertura y
cierre sinteticos -dos parejas y un grupo ambiguo- por comparar.py y generar_papel.py.
Y al final se mira el SKILL.md: huecos en vez de años fijos y el borrado con ruta literal.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

# La carpeta del script, explicita: lanzado con `python -I` no esta en sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib_continuidad import salida_utf8  # noqa: E402

# Los tres casos. El importe de la 710 se eligio; los demas son redondos a
# proposito para que el cuadre de control se lea de un vistazo: 383.333 tiene que
# ser 333.333 + 50.000 exactos.
CASOS = [
    ("710000000", "Prestacion de servicios",   0.0, 333333.0, -333333.0),
    ("570000000", "Caja, euros",          333333.0,      0.0,  333333.0),
    ("122000000", "Reserva nueva del cliente", 0.0,  50000.0,  -50000.0),
    ("572000001", "Banco c/c",             50000.0,      0.0,   50000.0),
]

ESPERADO = {
    "cuenta_nueva": {"122000000": -50000.0},
    "diferencia": {"570": 333333.0, "572": 50000.0},
    "sin_apertura": {},
}
CUADRE_ESPERADO = {
    "descuadre_apertura_balance": 383333.0,
    "por_cuentas_de_resultados": 333333.0,
    "por_cuentas_no_reconocidas": 50000.0,
    "sin_explicar": 0.0,
}


def _conectar(ruta: Path):
    import pyodbc

    return pyodbc.connect(
        "DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ=" + str(ruta) + ";",
        readonly=True,
    )


def consultar(ruta: Path, sql: str) -> list:
    """Lo que exportar_consulta(fuente='diario') escribiria: filas como dict, fechas en texto."""
    con = _conectar(ruta)
    try:
        cur = con.cursor()
        cur.execute(sql)
        cols = [d[0] for d in cur.description]
        filas = []
        for fila in cur.fetchall():
            filas.append({c: (str(v) if hasattr(v, "date") else
                              float(v) if v.__class__.__name__ == "Decimal" else v)
                          for c, v in zip(cols, fila)})
        return filas
    finally:
        con.close()


def columnas_diario(ruta: Path) -> list:
    con = _conectar(ruta)
    try:
        return [c.column_name for c in con.cursor().columns(table="Diario")]
    finally:
        con.close()


def amanar(origen: Path, destino: Path, asiento: str, concepto: str) -> None:
    """Copia el diario y le inserta los casos en el asiento de apertura."""
    import pyodbc

    shutil.copy2(origen, destino)
    con = pyodbc.connect(
        "DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ=" + str(destino) + ";"
    )
    cur = con.cursor()
    for cta, nom, debe, haber, ape in CASOS:
        cur.execute(
            "INSERT INTO Diario (FECHA, ASIENTO, CUENTA, NOMBRE, DEBE, HABER, "
            "CONCEPTO, APERTURA) VALUES (?,?,?,?,?,?,?,?)",
            "2025-01-01", asiento, cta, nom, debe, haber, concepto, ape,
        )
    con.commit()
    con.close()


def exportar(diario: Path, trabajo: Path, nombre: str) -> tuple:
    """Las dos exportaciones del SKILL.md, por ODBC. Devuelve (filtro.json, apertura.json, filtro)."""
    from lib_continuidad import SQL_FILTRO, leer_filtro, sql_apertura

    ruta_filtro = trabajo / (nombre + "_filtro.json")
    ruta_filtro.write_text(json.dumps(consultar(diario, SQL_FILTRO), ensure_ascii=False),
                           encoding="utf-8")
    filtro = leer_filtro(ruta_filtro)
    if filtro.get("ausente"):
        raise SystemExit("El diario de partida no sirve: " + filtro["motivo"])
    con_apertura = "APERTURA" in [c.upper() for c in columnas_diario(diario)]
    ruta_ap = trabajo / (nombre + "_apertura.json")
    ruta_ap.write_text(json.dumps(consultar(diario, sql_apertura(filtro, con_apertura)),
                                  ensure_ascii=False), encoding="utf-8")
    return ruta_filtro, ruta_ap, filtro


# Caso sintetico de diferencias que se compensan (registro del 21/10/2026: en un expediente
# real salieron dos parejas -una 40 contra una 57, y dos 43 entre si- que el auditor emparejo
# a mano). Cuenta de maximo nivel -> (saldo de cierre, saldo de apertura). Las cuentas de
# apertura son las de nivel con ceros detras, para que pasen por el mapeo de prefijo.
#   400/570      +1.000,00 / -1.000,00          pareja (la 40 contra la 57)
#   430/431      +250,40 / -250,40              pareja (dos 43 entre si)
#   520,521/523  +777,77 dos veces / -777,77    AMBIGUA: dos formas de emparejar
#   572/4750     +100,00 / -100,01              no: difieren en un centimo
#   4751         -777,76                        sola, cuadra la apertura a cero
#   1700/1710    sin_apertura -300 / +300       no: solo se emparejan las `diferencia`
SINTETICO = {
    "100": (-10000.0, -10000.0),
    "400": (-3000.0, -2000.0),
    "570": (1500.0, 500.0),
    "430": (5000.0, 5250.40),
    "431": (2000.0, 1749.60),
    "572": (4500.0, 4600.0),
    "520": (0.0, 777.77),
    "521": (0.0, 777.77),
    "523": (0.0, -777.77),
    "4750": (0.0, -100.01),
    "4751": (0.0, -777.76),
    "1700": (-300.0, None),
    "1710": (300.0, None),
}
COMPENSA_ESPERADO = [
    {"clase": "pareja", "importe": 1000.0, "positivas": ["400"], "negativas": ["570"]},
    {"clase": "ambigua", "importe": 777.77, "positivas": ["520", "521"], "negativas": ["523"]},
    {"clase": "pareja", "importe": 250.4, "positivas": ["430"], "negativas": ["431"]},
]
COLUMNA_G_ESPERADA = {"400": "con 570", "570": "con 400", "430": "con 431", "431": "con 430",
                      "520": "ambigua: con 523, igual que 521", "521": "ambigua: con 523, igual que 520",
                      "523": "ambigua: con 520 o 521"}


def _probar_compensadas(aqui: Path, trabajo: Path, comprobar) -> None:
    """Filtro, apertura y cierre sinteticos -sin diario real- por comparar.py y generar_papel.py."""
    sint = trabajo / "sintetico"
    sint.mkdir(parents=True, exist_ok=True)
    (sint / "filtro.json").write_text(json.dumps([{"Datos": "ASIENTO|Numero|Es igual a|1|||"}]),
                                      encoding="utf-8")
    apertura = []
    for cta, (_, ap) in SINTETICO.items():
        if ap is None:
            continue
        apertura.append({"FECHA": "2025-01-01", "ASIENTO": 1, "CUENTA": cta.ljust(9, "0"),
                         "NOMBRE": "CUENTA " + cta, "DEBE": max(ap, 0.0), "HABER": max(-ap, 0.0),
                         "POR_FILTRO": 1})
    (sint / "apertura.json").write_text(json.dumps(apertura), encoding="utf-8")
    # el cierre como lo deja exportar_consulta: texto con coma decimal
    cierre = [{"Cuenta": cta, "Nombre": "Cuenta " + cta,
               "SaldoAuditoria": format(sc, ".4f").replace(".", ",")}
              for cta, (sc, _) in SINTETICO.items()]
    (sint / "cierre.json").write_text(json.dumps(cierre), encoding="utf-8")

    salida = sint / "resultado.json"
    res = subprocess.run(
        [sys.executable, "-I", str(aqui / "comparar.py"), "--filtro", str(sint / "filtro.json"),
         "--apertura", str(sint / "apertura.json"), "--cierre", str(sint / "cierre.json"),
         "--salida", str(salida), "--ejercicio", "2025", "--ejercicio-anterior", "2024",
         "--cliente", "CLIENTE SINTETICO"],
        capture_output=True, text=True, encoding="utf-8",
    )
    comprobar("sintetico: comparar.py termina bien", res.returncode == 0)
    comprobar("sintetico: la pantalla dice que importe lleva cada tipo -la diferencia es apertura "
              "menos cierre, la sin apertura es el saldo al cierre- (registro del 28/10/2026)",
              "importe = apertura − cierre" in res.stdout
              and "importe = saldo al cierre auditado" in res.stdout)
    if res.returncode != 0:
        print(res.stdout + res.stderr)
        return
    datos = json.loads(salida.read_text(encoding="utf-8"))
    comprobar("sintetico: la apertura cuadra a cero", abs(datos["cuadre"]["sin_explicar"]) < 0.005
              and abs(datos["cuadre"]["descuadre_apertura_balance"]) < 0.005)
    comprobar("sintetico: dos parejas y un grupo ambiguo, y nada mas (salio "
              + str(datos["compensaciones"]) + ")", datos["compensaciones"] == COMPENSA_ESPERADO)
    tipos = {h["cuenta"]: h["tipo"] for h in datos["hallazgos"]}
    marcadas = {h["cuenta"] for h in datos["hallazgos"] if h.get("compensacion")}
    comprobar("sintetico: siguen siendo diferencias -10 `diferencia`, 2 `sin_apertura`-",
              sorted(tipos.values()).count("diferencia") == 10
              and sorted(tipos.values()).count("sin_apertura") == 2)
    comprobar("sintetico: ni el centimo de diferencia (572/4750), ni la sola (4751), ni las "
              "`sin_apertura` (1700/1710) se emparejan",
              not marcadas & {"572", "4750", "4751", "1700", "1710"})
    comprobar("sintetico: la pantalla dice las parejas, la ambigua y que siguen siendo diferencias",
              "COMPENSAN  2 pareja(s) y 1 grupo(s) ambiguo(s)" in res.stdout
              and "AMBIGUA  777.77" in res.stdout and "Siguen siendo diferencias" in res.stdout
              and "[se compensa con 570]" in res.stdout)

    papel = sint / "papel_sintetico.xlsx"
    res = subprocess.run(
        [sys.executable, "-I", str(aqui / "generar_papel.py"), "--resultado", str(salida),
         "--fecha-cierre", "31/12/24", "--generado", "21/10/2026", "--salida", str(papel)],
        capture_output=True, text=True, encoding="utf-8",
    )
    comprobar("sintetico: generar_papel.py termina bien", res.returncode == 0)
    if res.returncode != 0:
        print(res.stdout + res.stderr)
        return
    import openpyxl

    from lib_uso_ia import HOJA

    wb = openpyxl.load_workbook(papel)
    ws = wb["AG)02"]
    col_g = {str(ws.cell(f, 1).value): ws.cell(f, 7).value for f in range(7, ws.max_row + 1)
             if ws.cell(f, 7).value}
    comprobar("sintetico: la columna G del papel marca cada pareja y la ambigua (salio "
              + str(col_g) + ")", col_g == COLUMNA_G_ESPERADA
              and "posible reclasificación al abrir" in str(ws["G6"].value))
    textos = " ".join(str(c.value) for fila in wb[HOJA].iter_rows() for c in fila if c.value)
    comprobar("sintetico: la hoja «Uso de IA» lo dice como marca, sin explicarlo",
              "2 pareja(s) única(s) y 1 grupo(s) ambiguo(s)" in textos
              and "siguen contando como diferencias" in textos)


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    p.add_argument("--diario", required=True, help="un .smn real, no se modifica")
    p.add_argument("--cierre", required=True)
    p.add_argument("--trabajo", default="trabajo_prueba")
    # el cierre que exporta el MCP son filas crudas, sin ejercicios: hay que decirlos
    p.add_argument("--ejercicio", default="2025")
    p.add_argument("--ejercicio-anterior", default="2024")
    args = p.parse_args()

    aqui = Path(__file__).resolve().parent
    trabajo = Path(args.trabajo)
    trabajo.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(aqui))
    import pandas as pd

    from lib_continuidad import COLUMNAS_APERTURA, aplicar_filtro, cargar_apertura

    fallos = []

    def comprobar(desc, ok):
        print(("OK  " if ok else "FALLA  ") + desc)
        if not ok:
            fallos.append(desc)

    # 1 · La condicion JET selecciona lo mismo que el filtro en pandas, sobre el diario real
    real = Path(args.diario)
    _, ap_real, filtro = exportar(real, trabajo, "real")
    diario = pd.DataFrame(consultar(real, "SELECT * FROM Diario"))
    m_pandas = aplicar_filtro(diario, filtro)
    df_ap = cargar_apertura(ap_real)
    sel = df_ap[df_ap["POR_FILTRO"]]
    neto_pandas = round(float((diario.loc[m_pandas, "DEBE"] - diario.loc[m_pandas, "HABER"]).sum()), 2)
    neto_jet = round(float((sel["DEBE"] - sel["HABER"]).sum()), 2)
    comprobar("filtro " + filtro["campo"] + " " + filtro["operador"] + " " + repr(filtro["valor"])
              + ": JET y pandas seleccionan los mismos apuntes (" + str(int(m_pandas.sum())) + ")",
              int(m_pandas.sum()) == len(sel) and abs(neto_pandas - neto_jet) < 0.005)
    columnas = set(df_ap.columns)
    comprobar("la apertura exportada lleva solo las columnas de la prueba: "
              + ", ".join(sorted(columnas)),
              columnas <= set(COLUMNAS_APERTURA) | {"APERTURA", "POR_FILTRO"} and "CONCEPTO" not in columnas)
    comprobar("y es una fraccion del diario: " + str(len(df_ap)) + " de " + str(len(diario)) + " apuntes",
              len(df_ap) < len(diario))
    from lib_continuidad import SQL_FILTRO
    skill_md = (aqui.parent / "SKILL.md").read_text(encoding="utf-8")
    comprobar("el SKILL.md da la consulta del filtro igual que el script", SQL_FILTRO in skill_md)
    comprobar("y el SKILL.md no manda subir el diario (ni gzip ni mdbtools)",
              not any(t in skill_md for t in ("gzip", "mdb-export", "mdbtools", "diario.smn.gz")))

    # 2 · El fixture: el asiento y el concepto de apertura se leen del diario real, porque
    # los casos tienen que ir en el MISMO asiento o el filtro no los ve
    asiento = str(diario.loc[m_pandas, "ASIENTO"].iloc[0])
    concepto = str(diario.loc[m_pandas, "CONCEPTO"].iloc[0])
    print("diario de partida: asiento de apertura " + asiento + ", concepto " + repr(concepto))
    fixture = trabajo / "fixture.smn"
    amanar(real, fixture, asiento, concepto)
    ruta_filtro, ruta_ap, _ = exportar(fixture, trabajo, "fixture")

    def comparar(filtro_json, apertura_json, nombre):
        salida = trabajo / ("resultado_" + nombre + ".json")
        res = subprocess.run(
            [sys.executable, str(aqui / "comparar.py"), "--filtro", str(filtro_json),
             "--apertura", str(apertura_json), "--cierre", args.cierre, "--salida", str(salida),
             "--ejercicio", args.ejercicio, "--ejercicio-anterior", args.ejercicio_anterior,
             "--cliente", "CLIENTE DE PRUEBA"],
            capture_output=True, text=True,
        )
        if res.returncode != 0:
            print(res.stdout + res.stderr)
            raise SystemExit(2)
        return salida, json.loads(salida.read_text(encoding="utf-8"))

    def por_tipo(d):
        out: dict = {t: {} for t in ESPERADO}
        for h in d["hallazgos"]:
            out.setdefault(h["tipo"], {})[h["cuenta"]] = h["importe"]
        return out

    # Lo que el fixture AÑADE a la linea base del diario real: un expediente real puede traer
    # sus propias diferencias de apertura (el de calibracion de octubre trae seis), y el arnes
    # tiene que medir los tres casos metidos, no suponer un diario sin ninguna.
    _, base = comparar(trabajo / "real_filtro.json", ap_real, "real")
    salida, datos = comparar(ruta_filtro, ruta_ap, "fixture")
    b, f = por_tipo(base), por_tipo(datos)
    for tipo, esperado in ESPERADO.items():
        delta = {c: round(v - b.get(tipo, {}).get(c, 0.0), 2) for c, v in f.get(tipo, {}).items()
                 if abs(v - b.get(tipo, {}).get(c, 0.0)) >= 0.005}
        comprobar("hallazgos " + tipo + " que añade el fixture: " + str(esperado) + " (salio " + str(delta) + ")",
                  delta == esperado)
    print("    (la linea base del diario real trae " + str(len(base["hallazgos"])) + " hallazgos propios)")
    for clave, esperado in CUADRE_ESPERADO.items():
        real_v = round(datos["cuadre"][clave] - base["cuadre"][clave], 2)
        comprobar("cuadre " + clave + " += " + format(esperado, ".2f") + " (salio " + format(real_v, ".2f") + ")",
                  abs(real_v - esperado) <= 0.005)

    # 3 · El papel y su bloque de uso de IA
    papel = trabajo / "papel_fixture.xlsx"
    res = subprocess.run(
        [sys.executable, str(aqui / "generar_papel.py"), "--resultado", str(salida),
         "--fecha-cierre", "31/12/25", "--generado", "19/10/2026", "--salida", str(papel)],
        capture_output=True, text=True,
    )
    if res.returncode != 0:
        print(res.stdout + res.stderr)
        return 2
    import openpyxl

    from lib_uso_ia import EFECTO, HOJA, REVISION, VERIFICACION, VERIFICADO

    wb = openpyxl.load_workbook(papel)
    ws = wb[HOJA] if HOJA in wb.sheetnames else None
    comprobar("el papel lleva la hoja «" + HOJA + "», la primera", ws is not None and wb.sheetnames[0] == HOJA)
    if ws is not None:
        textos = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
        huecos = sorted({t.split("⟦IA:")[1].split(" ")[0] for t in textos if "⟦IA:" in t})
        comprobar("con los huecos que rellena la aplicacion: " + ", ".join(huecos),
                  huecos == ["aplicacion", "informacion", "modelo", "plugin", "usuario"] or
                  huecos == ["aplicacion", "entorno", "informacion", "modelo", "plugin", "usuario"])
        # lo del equipo, en blanco: la herramienta no escribe ni una letra ahi
        etiquetas = {e for e, _ in VERIFICACION + VERIFICADO + EFECTO + REVISION}
        llenas = [ws.cell(f, 1).value for f in range(1, ws.max_row + 1)
                  if ws.cell(f, 1).value in etiquetas and any(ws.cell(f, c).value for c in (2, 3, 4))]
        filas_equipo = sum(1 for f in range(1, ws.max_row + 1) if ws.cell(f, 1).value in etiquetas)
        comprobar("los " + str(filas_equipo) + " apartados del equipo salen en blanco",
                  not llenas and filas_equipo == len(VERIFICACION + VERIFICADO + EFECTO + REVISION))
        comprobar("y la hoja del papel dice que esta pendiente de verificacion",
                  "Pendiente de verificación" in str(wb["AG)02"]["H4"].value or ""))
    propia = aqui / "lib_uso_ia.py"
    otras = [p for p in aqui.parents[1].glob("*/scripts/lib_uso_ia.py") if p.resolve() != propia.resolve()]
    comprobar("lib_uso_ia.py es identico en los " + str(len(otras) + 1) + " skills que lo llevan",
              all(p.read_bytes() == propia.read_bytes() for p in otras))

    # La linea base real, solo informativa: un expediente real puede traer parejas propias
    # (el de calibracion de octubre trae una) y no se puede exigir ni un numero ni otro
    print("    (la linea base del diario real trae " + str(base["meta"].get("parejas_compensadas"))
          + " pareja(s) y " + str(base["meta"].get("grupos_compensacion_ambigua"))
          + " grupo(s) ambiguo(s) de diferencias que se compensan)")

    # 4 · Diferencias que se compensan, con un caso SINTETICO de punta a punta
    _probar_compensadas(aqui, trabajo, comprobar)

    # 5 · El SKILL.md: huecos en vez de años fijos, y el borrado con la ruta literal
    # (registro del 21/10/2026: el expediente probado era de 2024 y los ejemplos decian 2025;
    # y en Cowork la variable no llega a la orden de borrado y `$(cat …)` lo bloquea)
    import re

    fijos = re.findall(r"--ejercicio(?:-anterior)? \"?\d{4}|--fecha-cierre \"?\d", skill_md)
    comprobar("el SKILL.md no lleva años ni fechas fijos en los ejemplos (" + str(fijos) + ")",
              not fijos)
    borrados = re.findall(r"^rm -rf [^\n]*", skill_md, flags=re.M)
    comprobar("y borra el directorio de trabajo con la ruta literal: " + str(borrados),
              bool(borrados) and not any("$" in b for b in borrados))

    if fallos:
        print("\nFALLA: " + str(len(fallos)))
        return 1
    print("\nTodo detectado. El comparador ve los tres casos, el cuadre los explica, las "
          "diferencias que se compensan salen marcadas y el papel lleva el bloque de uso de IA "
          "con lo del equipo en blanco.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
