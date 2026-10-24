# -*- coding: utf-8 -*-
"""Arnes de estados-financieros. Comprobaciones PURAS: no toca ningun expediente.

    python probar_eeff.py

Lo que se prueba aqui es lo que se rompe en silencio: los signos de presentacion, que solo
sumen las lineas y las cuentas que deben, y que el orden sea el del modelo oficial y no el
alfabetico. Un fallo en cualquiera de esos tres da un papel con pinta correcta y cifras falsas.

Y los cuatro papeles de punta a punta, sobre un expediente SINTETICO que se escribe en una
carpeta temporal fuera del repositorio: cada uno lleva su hoja «Uso de IA» la primera, con los
huecos que rellena la aplicacion y los once apartados del equipo en blanco
(docs/evidencia-uso-ia.md), y lib_uso_ia.py es identico al de los demas skills.
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

# Los modulos de la carpeta se encuentran aunque se ejecute con `python -I` (20/10/2026).
sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib_eeff import (  # noqa: E402
    BORRADOR,
    SIN_HACER,
    ESTADO_BALANCE,
    lineas_pn,
    movimiento_cuenta,
    movimiento_por_epigrafe,
    ajustes_sin_epigrafe,
    ESTADO_PYG,
    a_float,
    cadena_del_saldo,
    cuadre_balance,
    cuentas_por_epigrafe,
    ejercicios,
    es_si,
    estructura_ordenada,
    hojas,
    salida_utf8,
    saldo,
)
from generar_ecpn import _pn  # noqa: E402
from generar_efe import ESTRUCTURALES, _agrupar, _fecha, _neto, _totales  # noqa: E402
from proponer_efe import (  # noqa: E402
    PREFIJO_IA,
    REGLAS,
    SIN_REGLA,
    Saldos,
    _hoja_catalogo,
    _r_amortizacion,
    _r_distribucion_beneficio,
    _r_gastos_financieros,
    _r_ingresos_financieros,
    agrupar,
    evaluar,
    proponer,
    regla_de,
    total_propuesto,
)
from generar_papel import _es_rotulo_pyg, _presentar_balance, _presentar_pyg, _sangria  # noqa: E402

# Los scripts que el skill ejecuta, en el orden de la secuencia. Todos tienen que arrancar con
# `python -I` desde otra carpeta: en Cowork la sesion los lanzo asi y no encontraban lib_eeff.
SCRIPTS = ("verificar_contrato.py", "generar_papel.py", "generar_ecpn.py", "generar_efe.py",
           "proponer_efe.py")


def _catalogo_de_prueba():
    """La hoja Catalogo de un expediente con dos plantillas: una con propuesta y otra sin."""
    from openpyxl import Workbook

    catalogo = [
        {"NumeroAsientoOA": "3", "Descripcion": "Amortizacion", "Aprobado": "False",
         "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
         "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
        for i, (cta, con) in enumerate(
            [("C68", "280"), ("280", "C68"), ("C68", "281"), ("281", "C68"),
             ("C68", "282"), ("282", "C68")], start=1)
    ] + [
        {"NumeroAsientoOA": "9", "Descripcion": "Lo que rellena el auditor", "Aprobado": "False",
         "Observaciones": "", "ApunteOA": "9", "Cuenta": "570", "Contrapartida": "129",
         "Concepto": "x", "Origen": "0", "Aplicacion": "0"},
    ]
    grupos = agrupar(catalogo)
    props = proponer(grupos, Saldos([{"Cuenta": "680", "SaldoAuditoria": "414,32"},
                                     {"Cuenta": "681", "SaldoAuditoria": "103.896,35"}]))
    wb = Workbook()
    wb.remove(wb.active)
    _hoja_catalogo(wb, grupos, {p["numero"]: total_propuesto(p) for p in props},
                   "CLIENTE DE PRUEBA", "2025")
    return wb["Catalogo"]


# ── los cuatro papeles de punta a punta, sobre un expediente sintetico ─────────
# Cifras inventadas pero coherentes entre si, para que ningun script pare: el balance cuadra
# en los dos ejercicios, la cadena del saldo cuadra en cada cuenta, el patrimonio neto abre en
# lo que cerro el anterior, el estado de flujos ata con el efectivo del balance y la propuesta
# de amortizacion coincide con el ajuste aprobado. Lo unico que NO cuadra es a proposito: el
# cliente no usa los grupos 8 y 9, y eso tiene que salir en el bloque de uso de IA con ATENCION.

GENERADO_PRUEBA = "19/10/2026"
HUECOS = ["aplicacion", "entorno", "informacion", "modelo", "plugin", "usuario"]


def _fixture(carpeta: Path) -> dict:
    def est(cod, estado, ap, clase, asig, s1, s2, concepto, mostrar="", **extra):
        f = {"CodigoCCAA": cod, "Estado": estado, "Activo_Pasivo": ap, "CodigoClase": clase,
             "CodigoAMostrar": mostrar, "Concepto": concepto, "CuentasAsignables": asig,
             "CodigoSeccion": "", "CodigoEpigrafe": "", "CodigoDesglose": "",
             "CodigoSubDesglose": "", "CodigoOrdenI": None,
             "SaldoAuditoria1": s1, "SaldoAuditoria2": s2, "SaldoAuditoria3": 0,
             "SaldoAuditoria4": 0, "SaldoAuditoria5": 0}
        f.update(extra)
        return f

    estructura = [
        est("Activo A)", "Balance", "ACTIVO", "A", "False", 600, 500, "A) ACTIVO NO CORRIENTE", "A)",
            CodigoSeccion="A"),
        est("Activo A)I.", "Balance", "ACTIVO", "A", "True", 600, 500, "I. Inmovilizado", "A)I.",
            CodigoSeccion="A", CodigoEpigrafe="1"),
        est("Activo B)VII.", "Balance", "ACTIVO", "A", "True", 450, 400,
            "VII. Efectivo y otros activos líquidos equivalentes", "B)VII.",
            CodigoSeccion="B", CodigoEpigrafe="7"),
        est("Pasivo A)", "Balance", "PASIVO", "P", "False", -900, -850, "A) PATRIMONIO NETO", "A)",
            CodigoSeccion="A"),
        est("Pasivo A)I.1.", "Balance", "PASIVO", "P", "True", -800, -800, "1. Capital escriturado",
            "A)I.1.", CodigoSeccion="A", CodigoEpigrafe="1", CodigoDesglose="1"),
        est("Pasivo A)I.7.", "Balance", "PASIVO", "P", "True", -100, -50,
            "VII. Resultado del ejercicio", "A)I.7.", CodigoSeccion="A", CodigoEpigrafe="7"),
        est("Pasivo C)V.", "Balance", "PASIVO", "P", "True", -150, -50,
            "V. Acreedores comerciales", "C)V.", CodigoSeccion="C", CodigoEpigrafe="5"),
        est("PyG 1.", "PyG", "", "I", "True", -500, -400, "1. Importe neto de la cifra de negocios",
            "1.", CodigoOrdenI="1", CodigoEpigrafe="1"),
        est("PyG 4.", "PyG", "", "I", "True", 380, 335, "4. Aprovisionamientos", "4.",
            CodigoOrdenI="4", CodigoEpigrafe="4"),
        est("PyG 8.", "PyG", "", "I", "True", 20, 15, "8. Amortización del inmovilizado", "8.",
            CodigoOrdenI="8", CodigoEpigrafe="8"),
        est("PyG A.5)", "PyG", "", "T", "False", -100, -50, "A.5) RESULTADO DEL EJERCICIO", "A.5)",
            CodigoOrdenI="22", CodigoSeccion="5"),
    ]

    def cta(cuenta, nombre, cod, asig, aud, ant, cli, aj=0.0, rec=0.0):
        return {"Cuenta": cuenta, "Nombre": nombre, "Digitos": str(len(cuenta)), "CodigoCCAA": cod,
                "AsignadaCCAA": asig, "SaldoAuditoria": aud, "SaldoAnterior": ant,
                "SaldoCliente": cli, "SaldoAj": aj, "SaldoRec": rec}

    cuentas = [
        cta("2100", "Construcciones", "Activo A)I.", "True", 700, 580, 700),
        cta("2810", "Amortización acumulada", "Activo A)I.", "True", -100, -80, -100),
        cta("5720", "Bancos", "Activo B)VII.", "True", 450, 400, 450),
        cta("1000", "Capital social", "Pasivo A)I.1.", "True", -800, -800, -800),
        cta("1290", "Resultado del ejercicio", "Pasivo A)I.7.", "True", -100, -50, -100),
        cta("129", "Resultado del ejercicio", "", "False", -100, -50, -100),
        cta("400", "Proveedores", "", "False", -150, -50, -140, aj=-10.0),
        cta("40000001", "PROV 40000001", "Pasivo C)V.", "True", -150, -50, -140, aj=-10.0),
        cta("7000", "Ventas", "PyG 1.", "True", -500, 0, -500),
        cta("6000", "Compras", "PyG 4.", "True", 380, 0, 370, aj=10.0),
        cta("680", "Amortización del inmovilizado", "PyG 8.", "True", 20, 0, 20),
    ]
    ejercicios_ = [{"CodigoAuditoria": "1", "Año": "2025"}, {"CodigoAuditoria": "2", "Año": "2024"}]

    def linea(orden, tipo, desc, o=0, ap=0, oa=0, aa=0):
        return {"CodigoAuditoria": "1", "CodigoOrden": str(orden), "CodigoLinea": tipo,
                "Descripcion": desc, "Origen": o, "Aplicacion": ap, "OrigenAjuste": oa,
                "AplicacionAjuste": aa, "OrigenAjustado": o + oa, "AplicacionAjustada": ap + aa,
                "SaldoAuditoriaAnterior": 0, "SaldoAuditoria": 0}

    cuadro = [
        linea(1, "F", "A) Flujos de efectivo de las actividades de explotación", o=220),
        linea(2, "C68", "Amortización del inmovilizado", oa=20),
        linea(3, "280", "Amortización acumulada", aa=20),
        linea(4, "F", "B) Flujos de efectivo de las actividades de inversión", ap=120),
        linea(5, "F", "C) Flujos de efectivo de las actividades de financiación", ap=50),
        linea(6, "T", "E) AUMENTO/DISMINUCIÓN NETA DEL EFECTIVO", ap=50),
    ]

    def apunte(n, desc, aprobado, i, cuenta, contra, o=0, ap=0, obs=""):
        return {"NumeroAsientoOA": str(n), "Fecha": "31/12/2025 0:00:00", "Descripcion": desc,
                "Aprobado": aprobado, "Observaciones": obs, "ApunteOA": str(i), "Cuenta": cuenta,
                "Concepto": "", "Contrapartida": contra, "Origen": o, "Aplicacion": ap}

    catalogo = [
        apunte(1, "Amortización", "True", 1, "C68", "280", o=20),
        apunte(1, "Amortización", "True", 2, "280", "C68", ap=20),
        apunte(1, "Amortización", "True", 3, "C68", "281"),
        apunte(1, "Amortización", "True", 4, "281", "C68"),
        apunte(1, "Amortización", "True", 5, "C68", "282"),
        apunte(1, "Amortización", "True", 6, "282", "C68"),
        apunte(2, "Distribución del resultado", "False", 1, "129", "120",
               obs="Apertura de la 129 repartida según el acuerdo de la Junta."),
        apunte(2, "Distribución del resultado", "False", 2, "11", "129"),
        apunte(2, "Distribución del resultado", "False", 3, "526", "129"),
        apunte(3, "Impuesto sobre beneficios", "False", 1, "G630", "C630",
               obs="Pagos por impuesto del ejercicio."),
        # con regla verificada y SIN importe: el plan no tiene la 66. Tiene que salir como «con
        # regla, no aplicada» y con el motivo, no entre las que no tienen regla (21/10/2026)
        apunte(4, "Gastos financieros", "False", 1, "C66", "G66", obs="Saldo de la 66."),
        apunte(4, "Gastos financieros", "False", 2, "G66", "C66"),
    ]
    # El caso normal en la practica: el cuadro de Gesia con cifras y NINGUN ajuste aprobado. El
    # estado esta SIN HACER y cuadra consigo mismo igual; no puede salir ni «NO CUADRA» ni verde.
    cuadro_sin = [linea(1, "F", "A) Flujos de efectivo de las actividades de explotación", o=220),
                  linea(4, "F", "B) Flujos de efectivo de las actividades de inversión", ap=120),
                  linea(5, "F", "C) Flujos de efectivo de las actividades de financiación", ap=50),
                  linea(6, "T", "E) AUMENTO/DISMINUCIÓN NETA DEL EFECTIVO", ap=50)]
    sin_aprobar = [dict(x, Aprobado="False", Origen=0, Aplicacion=0) for x in catalogo]
    rutas = {}
    for nombre, datos in (("estructura", estructura), ("cuentas", cuentas),
                          ("ejercicios", ejercicios_), ("cuadro", cuadro),
                          ("ajustes", catalogo), ("catalogo", catalogo),
                          ("cuadro_sin", cuadro_sin), ("sin_aprobar", sin_aprobar)):
        rutas[nombre] = carpeta / (nombre + ".json")
        rutas[nombre].write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    return rutas


def _papeles(carpeta: Path) -> dict:
    """Ejecuta los cuatro scripts como lo hace el skill. {script: (salida, codigo, ruta del papel)}."""
    r = _fixture(carpeta)
    aqui = Path(__file__).resolve().parent
    comun = ["--cliente", "CLIENTE DE PRUEBA", "--ejercicio", "2025", "--generado", GENERADO_PRUEBA]
    llamadas = {
        "generar_papel.py": ["--estructura", r["estructura"], "--cuentas", r["cuentas"],
                             "--ejercicios", r["ejercicios"]],
        "generar_ecpn.py": ["--estructura", r["estructura"], "--cuentas", r["cuentas"],
                            "--ejercicios", r["ejercicios"]],
        "generar_efe.py": ["--cuadro", r["cuadro"], "--ajustes", r["ajustes"],
                           "--estructura", r["estructura"], "--ejercicios", r["ejercicios"]],
        "proponer_efe.py": ["--catalogo", r["catalogo"], "--cuentas", r["cuentas"],
                            "--cuadro", r["cuadro"]],
        # el mismo papel de flujos, sin ningun ajuste aprobado
        "generar_efe.py sin aprobar": ["--cuadro", r["cuadro_sin"], "--ajustes", r["sin_aprobar"],
                                       "--estructura", r["estructura"],
                                       "--ejercicios", r["ejercicios"]],
    }
    res = {}
    for clave, args in llamadas.items():
        script = clave.split(" ")[0]
        papel = (carpeta / "AsistenteIA" / "EstadosFinancieros"
                 / (clave.replace(".py", "").replace(" ", "_") + ".xlsx"))
        # con -I y desde otra carpeta, que es como los lanza la sesion: si un script no
        # encuentra sus modulos, se ve aqui y no en casa del auditor
        p = subprocess.run([sys.executable, "-I", str(aqui / script)] + [str(x) for x in args]
                           + comun + ["--salida", str(papel)],
                           capture_output=True, text=True, encoding="utf-8", cwd=str(carpeta),
                           env=dict(os.environ, PYTHONIOENCODING="utf-8"))
        res[clave] = (p.stdout + p.stderr, p.returncode, papel)
    return res


def _bloque(papel: Path) -> dict:
    """Lo que hay que mirar de un papel: hojas, huecos, apartados del equipo, marca y fecha."""
    import openpyxl

    from lib_uso_ia import EFECTO, HOJA, MARCA_PAPEL, REVISION, VERIFICACION, VERIFICADO

    wb = openpyxl.load_workbook(papel)
    info = {"hojas": wb.sheetnames, "primera": wb.sheetnames[0] == HOJA}
    if HOJA not in wb.sheetnames:
        return info
    ws = wb[HOJA]
    textos = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
    info["huecos"] = sorted({t.split("⟦IA:")[1].split(" ")[0] for t in textos if "⟦IA:" in t})
    etiquetas = {e for e, _ in VERIFICACION + VERIFICADO + EFECTO + REVISION}
    filas = [f for f in range(1, ws.max_row + 1) if ws.cell(f, 1).value in etiquetas]
    info["equipo"] = (len(filas), sum(1 for f in filas if any(ws.cell(f, c).value for c in (2, 3, 4))))
    info["comprobaciones"] = [str(ws.cell(f, 2).value)[2:] for f in range(1, ws.max_row + 1)
                              if str(ws.cell(f, 2).value or "").startswith("· ")]
    info["resultado"] = next((str(ws.cell(f, 2).value) for f in range(1, ws.max_row + 1)
                              if ws.cell(f, 1).value == "Resultado"), "")
    datos = [n for n in wb.sheetnames if n != HOJA]
    info["marcadas"] = all(wb[n]["B4"].value == MARCA_PAPEL for n in datos)
    info["fecha"] = all(GENERADO_PRUEBA in str(wb[n]["A3"].value) for n in datos)
    # lo que dice la hoja Comprobaciones del propio papel, para contrastar el bloque con ella
    if "Comprobaciones" in wb.sheetnames:
        wc = wb["Comprobaciones"]
        info["etiquetas"] = [wc.cell(f, 2).value for f in range(1, wc.max_row + 1)
                             if wc.cell(f, 2).value not in (None, "Resultado")]
        info["no_cuadra"] = sum(1 for e in info["etiquetas"]
                                if e in ("NO CUADRA", "A TENER EN CUENTA", SIN_HACER, BORRADOR))
        info["cabecera"] = str(wc["A5"].value or "")
    return info


def _pruebas_uso_ia() -> list:
    """Los cuatro papeles, cada uno con su bloque de uso de IA. Fuera del repositorio."""
    from lib_eeff import frases_comprobacion

    pruebas = []
    with tempfile.TemporaryDirectory(prefix="probar_eeff_") as tmp:
        papeles = _papeles(Path(tmp))
        for script, (salida, codigo, papel) in papeles.items():
            if codigo not in (0, 1) or not papel.exists():
                pruebas.append((f"USO DE IA · {script} escribe su papel (salio {codigo}):\n{salida[-600:]}",
                                False))
                continue
            b = _bloque(papel)
            pruebas += [
                (f"USO DE IA · {script}: la hoja «Uso de IA» existe y es la primera", b["primera"]),
                (f"USO DE IA · {script}: los huecos de la aplicacion son exactamente "
                 + ", ".join(HUECOS) + f" (salieron {b.get('huecos')})", b.get("huecos") == HUECOS),
                (f"USO DE IA · {script}: los 11 apartados del equipo salen vacios "
                 f"(salieron {b.get('equipo')} filas/llenas)", b.get("equipo") == (11, 0)),
                (f"USO DE IA · {script}: cada hoja del papel lleva la marca de pendiente de verificacion",
                 b.get("marcadas", False)),
                (f"USO DE IA · {script}: la fecha es la de --generado, no la del reloj", b.get("fecha", False)),
                (f"USO DE IA · {script}: el bloque lleva comprobaciones de la herramienta",
                 len(b.get("comprobaciones", [])) >= 2),
            ]
            if script != "proponer_efe.py":
                # el bloque dice ATENCION exactamente donde la hoja Comprobaciones no cuadra
                atencion = sum(1 for c in b.get("comprobaciones", []) if c.startswith("ATENCIÓN"))
                pruebas.append((f"USO DE IA · {script}: el bloque dice ATENCIÓN tantas veces como la "
                                f"hoja Comprobaciones no cuadra ({atencion} y {b.get('no_cuadra')})",
                                atencion == b.get("no_cuadra")))
        b = _bloque(papeles["generar_ecpn.py"][2]) if papeles["generar_ecpn.py"][2].exists() else {}
        pruebas.append(("USO DE IA · ECPN: el cliente sin grupos 8 y 9 sale como ATENCIÓN en el bloque, "
                        "no como cumplido",
                        any(c.startswith("ATENCIÓN") and "grupos 8 y 9" in c
                            for c in b.get("comprobaciones", []))))
        b = _bloque(papeles["proponer_efe.py"][2]) if papeles["proponer_efe.py"][2].exists() else {}
        pruebas.append(("USO DE IA · PROPUESTA: el resultado dice que los importes los ha calculado "
                        "la herramienta, cuantos y que es un borrador",
                        "ha calculado los importes propuestos para 2 de las 4" in b.get("resultado", "")
                        and "BORRADOR" in b.get("resultado", "")))
        pruebas.append(("USO DE IA · PROPUESTA: el resultado nombra la plantilla con regla que no se "
                        "ha aplicado, aparte de las que no tienen regla",
                        "la nº 4 tiene regla verificada que no se ha aplicado" in b.get("resultado", "")))

        # ── tanda 3 (21/10/2026): cada regla dice si se aplico o por que no ──
        salida_p, _, papel_p = papeles["proponer_efe.py"]
        reglas_consola = salida_p.split("reglas verificadas", 1)[-1].split("PARA CONTAR", 1)[0]
        pruebas.append(("REGLAS · la consola lista las CUATRO reglas, una línea cada una: aplicada a "
                        "qué plantilla, o por qué no",
                        "reglas verificadas" in salida_p
                        and all(f"regla «{r}»:" in reglas_consola for r, _, _ in REGLAS)
                        and "regla «Amortización del ejercicio»: aplicada a la nº 1" in reglas_consola
                        and "regla «Gastos financieros»: NO aplicada a la nº 4, porque la 66 no está "
                            "en el plan del expediente" in reglas_consola
                        and "regla «Ingresos financieros»: NO aplicada, porque ninguna plantilla del "
                            "catálogo tiene su forma" in reglas_consola))
        entrega_p = [l.strip()[2:] for l in salida_p.split("PARA CONTAR", 1)[-1].splitlines()
                     if l.strip().startswith("- ")]
        pruebas.append(("REGLAS · lo que se cuenta al entregar separa «con regla, no aplicada» de "
                        f"«sin regla» (salio {entrega_p})",
                        any("reglas verificadas" in l and "no se han aplicado" in l
                            and "nº 4, porque la 66" in l for l in entrega_p)
                        and any("1 tiene regla verificada y no se ha aplicado" in l
                                and "1 no tiene regla" in l for l in entrega_p)))
        if papel_p.exists():
            import openpyxl
            wc = openpyxl.load_workbook(papel_p)["Catalogo"]
            filas = {str(wc.cell(f, 1).value): f for f in range(6, wc.max_row + 1)}
            col = lambda n, c: wc.cell(filas[n], c).value if n in filas else None  # noqa: E731
            pruebas.append(("CATALOGO · «Regla y motivo»: aplicada en la propuesta, el motivo "
                            "concreto en la que tiene regla y no se aplicó, y «sin regla» en las demás "
                            f"(salio {[(n, col(n, 7), col(n, 8)) for n in sorted(filas)]})",
                            wc.cell(5, 8).value == "Regla y motivo"
                            and col("1", 7) == "sí" and "aplicada" in str(col("1", 8))
                            and col("4", 7) == "no"
                            and col("4", 8) == "«Gastos financieros»: NO aplicada, porque la 66 no "
                                               "está en el plan del expediente"
                            and col("3", 7) == "—" and col("3", 8) == SIN_REGLA
                            and col("3", 9) == "Pagos por impuesto del ejercicio."))

        # ── tanda 3: cada linea de «PARA CONTAR AL ENTREGAR» se entiende leida sola ──
        salida_e = papeles["generar_ecpn.py"][0]
        pruebas.append(("ENTREGA · ECPN: la línea de los grupos 8 y 9 lleva su sujeto, no «NO los usa "
                        "el cliente» a secas (registro del 21/10/2026)",
                        "A TENER EN CUENTA: el cliente NO usa los grupos 8 y 9" in salida_e
                        and "NO los usa" not in salida_e))
        pruebas.append(("ENTREGA · EFE sin aprobar: la línea dice de qué son los ajustes que no están "
                        "aprobados",
                        "SIN HACER: NINGUNO de los 4 ajustes del catálogo del estado de flujos"
                        in papeles["generar_efe.py sin aprobar"][0]))

        # ── «NO CUADRA» solo para los fallos aritmeticos (prueba en Cowork, 20/10/2026) ──
        salida_p = papeles["proponer_efe.py"][0]
        pruebas.append(("ETIQUETAS · la propuesta es un BORRADOR, no un descuadre: ni la hoja "
                        "Comprobaciones ni lo que se cuenta al entregar dicen «NO CUADRA», y la "
                        f"cabecera no habla de descuadres (etiquetas {b.get('etiquetas')})",
                        BORRADOR in b.get("etiquetas", []) and "NO CUADRA" not in b.get("etiquetas", [])
                        and "NO CUADRA" not in salida_p and f"- {BORRADOR}: " in salida_p
                        and "BORRADOR" in b.get("cabecera", "")
                        and "DESCUADRE" not in b.get("cabecera", "")))
        salida_s, codigo_s, papel_s = papeles["generar_efe.py sin aprobar"]
        b = _bloque(papel_s) if papel_s.exists() else {}
        pruebas.append(("ETIQUETAS · el estado de flujos sin ningún ajuste aprobado está SIN HACER: "
                        "ni «NO CUADRA» ni «TODO CUADRA» en verde, y el bloque de uso de IA lo dice "
                        f"con ATENCIÓN (salio {codigo_s}, etiquetas {b.get('etiquetas')})",
                        codigo_s == 1 and SIN_HACER in b.get("etiquetas", [])
                        and "NO CUADRA" not in b.get("etiquetas", []) and "NO CUADRA" not in salida_s
                        and "SIN HACER" in b.get("cabecera", "")
                        and "TODO CUADRA" not in b.get("cabecera", "")
                        and any(c.startswith("ATENCIÓN, sin hacer") for c in b.get("comprobaciones", []))))
        pruebas.append(("ENTREGA · sin ajustes aprobados el script dice en una línea que el estado de "
                        "flujos está SIN HACER, para que el verde de lo demás no se lea como hecho; "
                        "con ajustes aprobados, no lo dice",
                        "estado de flujos: SIN HACER" in salida_s
                        and "estado de flujos: SIN HACER" not in papeles["generar_efe.py"][0]
                        and "estado de flujos: trabajado" in papeles["generar_efe.py"][0]))
    pruebas.append(("USO DE IA · una comprobacion que no se cumple no se repite en afirmativo: va "
                    "con ATENCIÓN y entre comillas, con el texto de lo que le pasa",
                    frases_comprobacion([("El balance cuadra", False, "difieren en 5,00 €"),
                                         ("Algo", None, "no se ha podido hacer"),
                                         ("Lo otro", True, "cuadra en las 3 cuentas"),
                                         ("Está trabajado", SIN_HACER, "ninguno aprobado"),
                                         ("Está terminado", BORRADOR, "faltan los grandes")])
                    == ["ATENCIÓN, no cuadra: «El balance cuadra» no se cumple. Difieren en 5,00 €.",
                        "ATENCIÓN, a tener en cuenta: «Algo» no se da por cumplida. No se ha podido hacer.",
                        "Lo otro. Cuadra en las 3 cuentas.",
                        "ATENCIÓN, sin hacer: «Está trabajado» no se cumple todavía. Ninguno aprobado.",
                        "ATENCIÓN, es un borrador: «Está terminado» no se cumple. Faltan los grandes."]))
    aqui = Path(__file__).resolve().parent
    with tempfile.TemporaryDirectory(prefix="probar_eeff_I_") as otra:
        arrancan = {s: subprocess.run([sys.executable, "-I", str(aqui / s), "--help"],
                                      capture_output=True, text=True, encoding="utf-8", cwd=otra,
                                      env=dict(os.environ, PYTHONIOENCODING="utf-8")).returncode
                    for s in SCRIPTS}
    pruebas.append(("ARRANQUE · los cinco scripts arrancan con `python -I` desde otra carpeta "
                    f"(salidas {arrancan})", all(c == 0 for c in arrancan.values())))
    pruebas.append(("USO DE IA · ningun script del skill lee el reloj",
                    all("date.today" not in (aqui / s).read_text(encoding="utf-8")
                        and "datetime.now" not in (aqui / s).read_text(encoding="utf-8")
                        for s in ("generar_papel.py", "generar_ecpn.py", "generar_efe.py",
                                  "proponer_efe.py"))))
    propia = aqui / "lib_uso_ia.py"
    otras = [p for p in aqui.parents[1].glob("*/scripts/lib_uso_ia.py") if p.resolve() != propia]
    pruebas.append((f"USO DE IA · lib_uso_ia.py es identico en los {len(otras) + 1} skills que lo llevan",
                    bool(otras) and all(p.read_bytes() == propia.read_bytes() for p in otras)))
    return pruebas


def _linea(codigo, estado, ap, clase, asignable, s1=0.0, s2=0.0, **extra):
    f = {"CodigoCCAA": codigo, "Estado": estado, "Activo_Pasivo": ap, "CodigoClase": clase,
         "CuentasAsignables": asignable, "SaldoAuditoria1": s1, "SaldoAuditoria2": s2,
         "Concepto": "x", "CodigoAMostrar": ""}
    f.update(extra)
    return f


def main() -> int:
    salida_utf8()
    fallos = 0

    # un balance de juguete: activo 1.000, pasivo -1.000 en crudo, mas un subtotal que
    # NO debe sumar (CuentasAsignables False)
    balance = [
        _linea("Activo A)", ESTADO_BALANCE, "ACTIVO", "A", "False", 1000.0, 900.0),
        _linea("Activo A)I.", ESTADO_BALANCE, "ACTIVO", "A", "True", 600.0, 500.0),
        _linea("Activo A)II.", ESTADO_BALANCE, "ACTIVO", "A", "True", 400.0, 400.0),
        _linea("Pasivo A)I.", ESTADO_BALANCE, "PASIVO", "P", "True", -1000.0, -900.0),
    ]
    pyg = [
        _linea("I.", ESTADO_PYG, "", "I", "True", -500.0, -400.0),      # ingreso, en crudo negativo
        _linea("II.", ESTADO_PYG, "", "I", "True", 120.0, 100.0),       # gasto, en crudo positivo
        _linea("A)", ESTADO_PYG, "", "ST", "False", 380.0, 300.0),      # subtotal positivo
    ]
    cuentas = [
        {"Cuenta": "7000", "CodigoCCAA": "I.", "AsignadaCCAA": "True", "SaldoAuditoria": -500.0,
         "SaldoCliente": -450.0, "SaldoAj": -50.0, "SaldoRec": 0.0, "Nombre": "Ventas"},
        {"Cuenta": "700", "CodigoCCAA": "I.", "AsignadaCCAA": "False", "SaldoAuditoria": -500.0,
         "SaldoCliente": -450.0, "SaldoAj": -50.0, "SaldoRec": 0.0, "Nombre": "agregada, NO suma"},
        {"Cuenta": "6000", "CodigoCCAA": "II.", "AsignadaCCAA": "True", "SaldoAuditoria": 120.0,
         "SaldoCliente": 120.0, "SaldoAj": 0.0, "SaldoRec": 0.0, "Nombre": "Compras"},
    ]

    pruebas = [
        ("importes: coma decimal española, punto de millar, y el formato ingles",
         a_float("1.234,56") == 1234.56 and a_float("1234.56") == 1234.56
         and a_float("") == 0.0 and a_float(None) == 0.0 and a_float("-1.000,00") == -1000.0),

        ("los booleanos del API llegan como texto",
         es_si("True") and es_si("-1") and es_si("Sí") and not es_si("False") and not es_si("")),

        ("solo son hojas las lineas que admiten cuentas: el subtotal queda fuera",
         len(hojas(balance)) == 3 and len(hojas(pyg)) == 2),

        ("EL CUADRE: activo mas pasivo suma cero sobre las hojas, en los dos ejercicios",
         abs(cuadre_balance(balance, 1)) < 0.01 and abs(cuadre_balance(balance, 2)) < 0.01),

        ("un balance descuadrado NO se da por bueno",
         abs(cuadre_balance(balance + [_linea("Pasivo A)II.", ESTADO_BALANCE, "PASIVO", "P", "True", -5.0)], 1)) > 0.01),

        ("presentacion del balance: el activo tal cual y el pasivo cambiado de signo",
         _presentar_balance(balance[1], 1) == 600.0 and _presentar_balance(balance[3], 1) == 1000.0),

        ("presentacion de la PyG, regla de Gesia: el negativo a magnitud, el positivo de epigrafe igual",
         _presentar_pyg(pyg[0], 1) == 500.0 and _presentar_pyg(pyg[1], 1) == 120.0),

        ("presentacion de la PyG: el subtotal POSITIVO se invierte",
         _presentar_pyg(pyg[2], 1) == -380.0),

        ("agregacion por epigrafe: solo cuentan las asignadas, la agregada no duplica",
         cuentas_por_epigrafe(cuentas, "SaldoAuditoria") == {"I.": -500.0, "II.": 120.0}),

        ("la cadena del saldo: cliente + ajustes + reclasificaciones = auditoria",
         (lambda c: abs((c[0] + c[1] + c[2]) - c[3]) < 0.01)(cadena_del_saldo(cuentas[0]))),

        ("los ejercicios se leen de Auditorias y el 1 no se supone que sea un año",
         ejercicios([{"CodigoAuditoria": "1", "Año": "2025"}, {"CodigoAuditoria": "2", "Año": "2024"}])
         == {"1": "2025", "2": "2024"}),

        ("orden del modelo, no alfabetico: el epigrafe 2 va antes que el 10",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("b", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="10"),
             _linea("a", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="2")])]
         == ["a", "b"]),

        ("LAS SECCIONES NO SE MEZCLAN: el no corriente entero antes que el corriente, aunque "
         "la seccion sea una LETRA y el epigrafe del corriente sea menor",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("B.1", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="B", CodigoEpigrafe="1"),
             _linea("A.2", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="2"),
             _linea("A.hd", ESTADO_BALANCE, "ACTIVO", "A", "False", CodigoSeccion="A")])]
         == ["A.hd", "A.2", "B.1"]),

        ("la PyG se ordena por CodigoOrdenI, y dentro de un mismo orden el detalle va antes "
         "que el subtotal: «20. Impuestos» antes de «A.4) RESULTADO...»",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("A.4)", ESTADO_PYG, "", "T", "False", CodigoOrdenI="20", CodigoSeccion="4"),
             _linea("21.", ESTADO_PYG, "", "I", "True", CodigoOrdenI="21", CodigoEpigrafe="21"),
             _linea("20.", ESTADO_PYG, "", "I", "True", CodigoOrdenI="20", CodigoEpigrafe="20")])]
         == ["20.", "A.4)", "21."]),

        ("ningun ajuste se pierde: el de 4100 llega por 410, el grupo 0 de cuadre de Gesia no "
         "cuenta, y una cuenta ajustada sin epigrafe ni agregada SI se denuncia",
         [c for c, _, _ in ajustes_sin_epigrafe([
             {"Cuenta": "4100", "AsignadaCCAA": "False", "SaldoAj": -53.07, "SaldoRec": 0.0, "Nombre": ""},
             {"Cuenta": "410", "AsignadaCCAA": "True", "SaldoAj": -53.07, "SaldoRec": 0.0, "Nombre": ""},
             {"Cuenta": "080", "AsignadaCCAA": "False", "SaldoAj": 999.0, "SaldoRec": 0.0, "Nombre": "Explotacion"},
             {"Cuenta": "8990", "AsignadaCCAA": "False", "SaldoAj": 12.0, "SaldoRec": 0.0, "Nombre": "huerfana"},
         ])] == ["8990"]),

        ("el balance va antes que la PyG, y dentro del balance el activo antes que el pasivo",
         [f["Estado"] for f in estructura_ordenada(balance + pyg)][:4] == [ESTADO_BALANCE] * 4),

        ("los rotulos de bloque de la PyG van SIN cifra: «B OPERACIONES INTERRUMPIDAS» guarda "
         "el importe de las continuadas y escribirlo seria inventar un resultado",
         _es_rotulo_pyg(_linea("B", ESTADO_PYG, "", "I", "False", 5432881.57, CodigoEpigrafe="0"), ESTADO_PYG)
         and not _es_rotulo_pyg(_linea("1.", ESTADO_PYG, "", "I", "True", 100.0, CodigoEpigrafe="1"), ESTADO_PYG)
         and not _es_rotulo_pyg(_linea("A)", ESTADO_BALANCE, "ACTIVO", "A", "False", 100.0, CodigoEpigrafe="0"), ESTADO_BALANCE)),

        ("la sangria crece con el nivel del epigrafe, y no se pasa de tres",
         _sangria(_linea("x", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A")) == 0
         and _sangria(_linea("x", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A",
                             CodigoEpigrafe="1", CodigoDesglose="1", CodigoSubDesglose="1")) == 3),

        ("el patrimonio neto son las lineas del PASIVO de la seccion A, y se reconoce por ahi "
         "y no por el texto del rotulo, que cambia entre modelos",
         [f["CodigoCCAA"] for f in lineas_pn([
             _linea("PN", ESTADO_BALANCE, "PASIVO", "P", "True", CodigoSeccion="A"),
             _linea("exigible", ESTADO_BALANCE, "PASIVO", "P", "True", CodigoSeccion="B"),
             _linea("activo", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A"),
             _linea("pyg", ESTADO_PYG, "", "I", "True", CodigoSeccion="A")])] == ["PN"]),

        ("EL MOVIMIENTO DEL EJERCICIO es cierre menos apertura, NO el Debe menos el Haber: "
         "esos acumulados llevan dentro el saldo de apertura",
         movimiento_cuenta({"SaldoAnterior": "-2.328.453,48", "DebeCliente": "0",
                            "HaberCliente": "2.876.858,37", "SaldoCliente": "-2.876.858,37",
                            "SaldoAj": "0", "SaldoRec": "0",
                            "SaldoAuditoria": "-2.876.858,37"})[1] == -548404.89),

        ("apertura + movimiento + ajustes + reclasificaciones = cierre, y solo suman las "
         "cuentas asignadas al componente",
         (lambda m: abs(sum(m["Pasivo A)I.3."][:4]) - m["Pasivo A)I.3."][4]) < 0.01
                 and abs(m["Pasivo A)I.3."][0] - (-2328453.48)) < 0.01)(
             movimiento_por_epigrafe([
                 {"Cuenta": "113", "CodigoCCAA": "Pasivo A)I.3.", "AsignadaCCAA": "True",
                  "SaldoAnterior": -2328453.48, "SaldoCliente": -2876858.37, "SaldoAj": 145327.30,
                  "SaldoRec": 0.0, "SaldoAuditoria": -2731531.07},
                 {"Cuenta": "11", "CodigoCCAA": "Pasivo A)I.3.", "AsignadaCCAA": "False",
                  "SaldoAnterior": -2328453.48, "SaldoCliente": -2876858.37, "SaldoAj": 145327.30,
                  "SaldoRec": 0.0, "SaldoAuditoria": -2731531.07},
             ], {"Pasivo A)I.3."}))),

        ("el patrimonio neto se presenta en positivo, y uno NEGATIVO sigue saliendo negativo",
         _pn(-61404.0) == 61404.0 and _pn(867.04) == -867.04),

        ("EFE · el ajuste es partida doble ENTRE LINEAS DEL ESTADO, y sus cobros tienen que "
         "igualar a sus pagos: un descuadre se ve",
         _totales([{"Origen": "100,00", "Aplicacion": "0"},
                   {"Origen": "0", "Aplicacion": "100,00"}])[2] == 0.0
         and _totales([{"Origen": "100,00", "Aplicacion": "0"}])[2] == 100.0),

        ("EFE · los ajustes llegan de un LEFT JOIN, con la cabecera repetida en cada apunte y "
         "la plantilla sin usar sin ninguno",
         (lambda g: len(g) == 2 and len(g["1"]["apuntes"]) == 2 and g["1"]["aprobado"]
                 and g["7"]["apuntes"] == [] and not g["7"]["aprobado"])(
             _agrupar([
                 {"NumeroAsientoOA": "1", "Descripcion": "Amortizacion", "Aprobado": "True",
                  "ApunteOA": "1", "Origen": "10", "Aplicacion": "0"},
                 {"NumeroAsientoOA": "1", "Descripcion": "Amortizacion", "Aprobado": "True",
                  "ApunteOA": "2", "Origen": "0", "Aplicacion": "10"},
                 {"NumeroAsientoOA": "7", "Descripcion": "Plantilla sin usar", "Aprobado": "False",
                  "ApunteOA": "", "Origen": "", "Aplicacion": ""},
             ]))),

        ("EFE · el estado se lee en neto, cobro menos pago",
         _neto({"OrigenAjustado": "1.541.971,78", "AplicacionAjustada": "0"},
               "OrigenAjustado", "AplicacionAjustada") == 1541971.78
         and _neto({"OrigenAjustado": "0", "AplicacionAjustada": "92.933,78"},
                   "OrigenAjustado", "AplicacionAjustada") == -92933.78),

        ("EFE · las lineas estructurales del cuadro se reconocen por CodigoLinea, que en las de "
         "detalle lleva la cuenta o el concepto",
         all(t in ESTRUCTURALES for t in ("F", "T", "S", "P"))
         and not any(t in ESTRUCTURALES for t in ("129", "C68", "G630", "57"))),

        ("EFE · la fecha viene con la hora pegada y cortarla a diez deja una fecha que no existe",
         _fecha("31/12/2025 0:00:00") == "31/12/2025" and _fecha("31/12/2025") == "31/12/2025"
         and _fecha(None) == ""),

        ("PROPUESTA · la plantilla se reconoce por LA FORMA de sus apuntes, no por el texto de "
         "la descripcion, que el auditor puede cambiar",
         regla_de({"apuntes": [{"Cuenta": "C68", "Contrapartida": "280"},
                               {"Cuenta": "C68", "Contrapartida": "281"},
                               {"Cuenta": "C68", "Contrapartida": "282"}]})[0]
         == "Amortización del ejercicio"
         and regla_de({"apuntes": [{"Cuenta": "999", "Contrapartida": "998"}]})[0] is None),

        ("PROPUESTA · los saldos se leen distinto segun el grupo: la de resultados por su SALDO "
         "del año, la de balance por su MOVIMIENTO",
         (lambda s: s.importe("66") == 220875.62 and s.importe("113") == 548404.89)(
             Saldos([{"Cuenta": "66", "SaldoAuditoria": "220.875,62", "SaldoAnterior": "0",
                      "SaldoCliente": "74.629,14"},
                     {"Cuenta": "113", "SaldoAuditoria": "-2.731.531,07",
                      "SaldoAnterior": "-2.328.453,48", "SaldoCliente": "-2.876.858,37"}]))),

        ("PROPUESTA · una cuenta que NO esta en el plan da None y no cero: proponer cero seria "
         "inventarse que el concepto no se ha movido",
         Saldos([]).importe("680") is None and Saldos([]).apertura("129") is None),

        ("PROPUESTA · con el resultado anterior en PERDIDA la regla de beneficios no propone "
         "nada, que es lo correcto mientras la de perdidas no se mida",
         proponer(agrupar([
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "1", "Cuenta": "129", "Contrapartida": "120",
              "Concepto": "(C)", "Origen": "0", "Aplicacion": "0"},
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "2", "Cuenta": "11", "Contrapartida": "129",
              "Concepto": "(P)", "Origen": "0", "Aplicacion": "0"},
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "3", "Cuenta": "526", "Contrapartida": "129",
              "Concepto": "(P)", "Origen": "0", "Aplicacion": "0"},
         ]), Saldos([{"Cuenta": "129", "SaldoAnterior": "5.176.303,90",
                      "SaldoCliente": "5.432.881,57", "SaldoAuditoria": "5.432.881,57"}])) == []),

        ("PROPUESTA · la amortizacion empareja cada dotacion con SU acumulada, el tramo sin "
         "dotacion se queda A CERO en vez de desaparecer, y la propuesta cuadra",
         (lambda ps: len(ps) == 1
                 and dict(((l, c), (co, pa)) for (l, c), _, co, pa in ps[0]["lineas"])
                     == {("C68", "280"): (414.32, 0.0), ("280", "C68"): (0.0, 414.32),
                         ("C68", "281"): (103896.35, 0.0), ("281", "C68"): (0.0, 103896.35),
                         ("C68", "282"): (0.0, 0.0), ("282", "C68"): (0.0, 0.0)}
                 and abs(sum(co for _, _, co, _ in ps[0]["lineas"])
                         - sum(pa for _, _, _, pa in ps[0]["lineas"])) < 0.01)(
             proponer(agrupar([
                 {"NumeroAsientoOA": "3", "Descripcion": "Amortizacion", "Aprobado": "False",
                  "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
                  "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
                 for i, (cta, con) in enumerate(
                     [("C68", "280"), ("280", "C68"), ("C68", "281"), ("281", "C68"),
                      ("C68", "282"), ("282", "C68")], start=1)
             ]), Saldos([{"Cuenta": "680", "SaldoAuditoria": "414,32"},
                         {"Cuenta": "681", "SaldoAuditoria": "103.896,35"}])))),

        ("PROPUESTA · lo que se propone se firma: el prefijo «Asistente IA:» va en la "
         "descripcion, para que no se confunda nunca con el trabajo del auditor",
         PREFIJO_IA.startswith("Asistente IA:")),

        # ── tanda 3 (21/10/2026): una regla que no da importe dice POR QUE ──
        # En un expediente real la distribucion y los gastos financieros quedaron en «—» sin
        # motivo y la hoja Comprobaciones los conto entre los que no tienen regla.
        ("MOTIVO · la 66 o la 76 que no estan en el plan no es lo mismo que estar a cero",
         _r_gastos_financieros([], Saldos([])) == (None, "la 66 no está en el plan del expediente")
         and _r_gastos_financieros([], Saldos([{"Cuenta": "66", "SaldoAuditoria": "0"}]))
         == (None, "la 66 no tiene saldo en el ejercicio (0,00)")
         and _r_ingresos_financieros([], Saldos([]))[1] == "la 76 no está en el plan del expediente"),

        ("MOTIVO · la amortizacion sin dotaciones dice que le pasa a cada una de las tres",
         (lambda r: r[0] is None and "680 no tiene saldo" in r[1] and "681 no está en el plan" in r[1]
                 and "682 no está en el plan" in r[1])(
             _r_amortizacion([], Saldos([{"Cuenta": "680", "SaldoAuditoria": "0"}])))),

        ("MOTIVO · la distribucion dice si la 129 falta, abre a cero o abre deudora (perdida, "
         "variante no medida)",
         "no está en el plan" in _r_distribucion_beneficio([], Saldos([]))[1]
         and "es cero" in _r_distribucion_beneficio([], Saldos([{"Cuenta": "129", "SaldoAnterior": "0"}]))[1]
         and (lambda m: "deudora" in m and "pérdida" in m and "no está contrastada" in m)(
             _r_distribucion_beneficio([], Saldos([{"Cuenta": "129", "SaldoAnterior": "5.176.303,90"}]))[1])),

        ("MOTIVO · la plantilla con la forma de una regla y sin importe queda como «con regla, no "
         "aplicada» con su motivo, y la regla lo dice en una frase que se entiende sola",
         (lambda r: r[0] == [] and r[1] == {"12": ("Gastos financieros",
                                                   "la 66 no está en el plan del expediente")}
                 and ("Gastos financieros", False,
                      "regla «Gastos financieros»: NO aplicada a la nº 12, porque la 66 no está en "
                      "el plan del expediente") in r[2])(
             evaluar(agrupar([
                 {"NumeroAsientoOA": "12", "Descripcion": "Gastos financieros", "Aprobado": "False",
                  "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
                  "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
                 for i, (cta, con) in enumerate([("C66", "G66"), ("G66", "C66")], start=1)]),
                 Saldos([])))),

        ("MOTIVO · sin ninguna plantilla con la forma de la regla, lo dice, y nombra la más "
         "parecida con el apunte que le falta, sin proponerla",
         (lambda r: r[0] == []
                 and r[1].get("1") == ("Distribución del resultado (beneficio)",
                                       "no tiene la forma de la regla: le falta el apunte 526 → 129")
                 and any(not ap and "ninguna plantilla del catálogo tiene su forma" in f
                         and "la más parecida es la nº 1" in f for _, ap, f in r[2]))(
             evaluar(agrupar([
                 {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
                  "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
                  "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
                 for i, (cta, con) in enumerate([("129", "120"), ("11", "129")], start=1)]),
                 Saldos([{"Cuenta": "129", "SaldoAnterior": "-100", "SaldoCliente": "-80",
                          "SaldoAuditoria": "-80"}])))),

        ("MOTIVO · los importes no cambian por llevar motivo: la amortizacion aplicada sigue "
         "proponiendo lo mismo y la regla sale como aplicada",
         (lambda r: len(r[0]) == 1 and all(ap for n, ap, _ in r[2] if n == "Amortización del ejercicio")
                 and total_propuesto(r[0][0]) == 104310.67)(
             evaluar(agrupar([
                 {"NumeroAsientoOA": "3", "Descripcion": "Amortizacion", "Aprobado": "False",
                  "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
                  "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
                 for i, (cta, con) in enumerate(
                     [("C68", "280"), ("280", "C68"), ("C68", "281"), ("281", "C68"),
                      ("C68", "282"), ("282", "C68")], start=1)]),
                 Saldos([{"Cuenta": "680", "SaldoAuditoria": "414,32"},
                         {"Cuenta": "681", "SaldoAuditoria": "103.896,35"}])))),

        ("ENTREGA · ninguna linea de «PARA CONTAR AL ENTREGAR» queda sin sujeto: los detalles "
         "a secas -«difieren en», «NO los usa el cliente», «N descuadran:»- no vuelven (21/10/2026)",
         not any(t in io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), g),
                              encoding="utf-8").read()
                 for g in ("generar_efe.py", "generar_papel.py", "generar_ecpn.py", "proponer_efe.py")
                 for t in ('else f"difieren en', '"NO los usa', ')} descuadran: "',
                           ')} no coinciden: "', ')} líneas no coinciden: "', ')} no cuadran: "',
                           ')} cuentas no cuadran:', ')} epígrafes no coinciden: {'))),

        ("ENTREGA · la linea que se cuenta al auditor NO repite el titulo de la comprobacion: "
         "el titulo es la afirmacion que se comprueba, y pegado a su detalle se leia al reves "
         "-«El estado esta trabajado ... NINGUNO esta aprobado»- (25/09/2026)",
         all('{t} \u00b7 {d}' not in io.open(
                 os.path.join(os.path.dirname(os.path.abspath(__file__)), g), encoding="utf-8").read()
             for g in ("generar_efe.py", "generar_papel.py", "generar_ecpn.py", "proponer_efe.py"))),

        ("ENTREGA · y el aviso de que no hay ningun ajuste aprobado sigue diciendo que el cuadro "
         "cuadra consigo mismo aunque el estado no este hecho, que es lo que enganya",
         (lambda t: "no es" in t and "todav" in t and "verde de las dem" in t)(
             io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "generar_efe.py"),
                     encoding="utf-8").read())),

        ("CATALOGO · el importe propuesto se ve en la hoja Catalogo, al lado del ya aprobado y "
         "con otro nombre. Los dos compartian la columna «Importe puesto», que suma lo YA "
         "REGISTRADO en el expediente -0,00 mientras nada esta aprobado-, y se leyo como que la "
         "propuesta venia vacia aunque la hoja Propuestas traia las cifras (26/09/2026)",
         (lambda ws: [ws.cell(row=5, column=c).value for c in (5, 6)]
          == ["Importe ya aprobado en Gesia", "Importe propuesto"]
          and ws.cell(row=6, column=5).value == 0.0
          and ws.cell(row=6, column=6).value == 104310.67)(_catalogo_de_prueba())),

        ("CATALOGO · y la plantilla SIN propuesta deja la celda en blanco, no a 0,00: un cero "
         "ahi es justo la lectura que se quiere evitar",
         _catalogo_de_prueba().cell(row=7, column=6).value is None),

        ("REDONDEO CONTABLE: el medio centimo sube, y no se redondea al par como hace round()",
         saldo({"SaldoAuditoria1": "1.000,555"}, 1) == 1000.56
         and saldo({"SaldoAuditoria1": "0,125"}, 1) == 0.13
         and saldo({"SaldoAuditoria1": "0,135"}, 1) == 0.14
         and round(0.125, 2) == 0.12),
    ]

    # ── 2.0.7: los nombres de terceros tienen que venir tokenizados ─────────
    from lib_eeff import exigir_tokens, nombres_en_claro
    plan = [{"Cuenta": "40000012", "Nombre": "PROV 40000012"},      # tokenizado: bien
            {"Cuenta": "43000001", "Nombre": "CLIENTE FICTICIO, S.A."},  # en claro: mal
            {"Cuenta": "4100", "Nombre": ""},                         # vacio: bien
            {"Cuenta": "400", "Nombre": "Proveedores"},               # agregada, no es de tercero
            {"Cuenta": "70000001", "Nombre": "Ventas"}]               # no es de tercero
    aqui = os.path.dirname(os.path.abspath(__file__))
    fuentes = {n: io.open(os.path.join(aqui, n), encoding="utf-8").read()
               for n in ("verificar_contrato.py", "generar_papel.py", "generar_ecpn.py", "proponer_efe.py")}
    salida = io.StringIO()
    _stdout, sys.stdout = sys.stdout, salida
    try:
        para = exigir_tokens(plan)
        sigue_con_permiso = exigir_tokens(plan, en_claro_por_el_auditor=True)
    finally:
        sys.stdout = _stdout
    pruebas += [
        ("TOKENS · solo cuenta como en claro la cuenta de TERCERO a maximo detalle con nombre que no es token",
         nombres_en_claro(plan) == ["43000001"]),
        ("TOKENS · con un nombre en claro se para, y lo dice con el numero de cuenta, nunca con el nombre",
         para is False and "43000001" in salida.getvalue() and "FICTICIO" not in salida.getvalue()),
        ("TOKENS · solo sigue si el auditor ha apagado la seudonimizacion (--nombres-en-claro)",
         sigue_con_permiso is True),
        ("TOKENS · los cuatro scripts que leen el plan lo comprueban antes de generar nada",
         all("exigir_tokens(cuentas, a.nombres_en_claro)" in s and "--nombres-en-claro" in s
             for s in fuentes.values())),
    ]

    # ── evidencia del uso de IA (docs/evidencia-uso-ia.md), en los cuatro papeles ──
    pruebas += _pruebas_uso_ia()

    for descripcion, ok in pruebas:
        print(f"  {'OK   ' if ok else 'FALLA'}  {descripcion}")
        if not ok:
            fallos += 1

    print()
    print("RESULTADO:", "todo correcto" if fallos == 0 else f"{fallos} fallo(s)")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
