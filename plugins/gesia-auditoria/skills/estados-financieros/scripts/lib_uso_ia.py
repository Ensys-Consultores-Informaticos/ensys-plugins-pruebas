"""El bloque «Utilizacion de herramientas de inteligencia artificial» del papel de trabajo.

Es la evidencia del uso de IA que queda en el expediente (docs/evidencia-uso-ia.md,
decidido el 19/10/2026): con que herramienta se hizo, que se le pidio, que informacion
se le facilito y como, que devolvio y que comprobo ella misma. Y deja EN BLANCO lo que
es del equipo: la verificacion, quien la hizo, el efecto en la conclusion y la revision.

El script escribe el bloque entero. Lo que solo sabe la aplicacion -en que aplicacion
corre, que plugin y version, quien lo ejecuto y que datos salieron del equipo- va como
hueco ⟦IA:clave · …⟧, y lo rellena la herramienta sellar_uso_ia del MCP sobre el papel
ya archivado en el disco del auditor. Si nadie lo sella, el hueco se lee tal cual y dice
que falta: un papel sin sellar no pasa por sellado.

El contenido se define UNA vez (contenido()) y se escribe en Excel (hoja_uso_ia) o en
Word (apartado_uso_ia): los dos formatos dicen lo mismo con las mismas palabras.

ESTE FICHERO ES EL MISMO EN TODOS LOS SKILLS QUE ESCRIBEN PAPEL. No se edita en uno solo:
se cambia en todos a la vez, y el arnes de cada skill compara su copia con las demas.

Nada aqui lee el reloj: la fecha de generacion llega como argumento (regla 4).
"""

from __future__ import annotations

import re

TITULO ="UTILIZACIÓN DE HERRAMIENTAS DE INTELIGENCIA ARTIFICIAL"
HOJA = "Uso de IA"
PENDIENTE = ("PENDIENTE DE VERIFICACIÓN POR EL EQUIPO. Mientras los apartados 5 a 7 estén en "
             "blanco, el resultado de este papel no es evidencia de auditoría ni conclusión del "
             "auditor.")
MARCA_PAPEL = "Pendiente de verificación: ver «" + HOJA + "»"
MARCA_DOCX = ("Pendiente de verificación: ver el apartado «Utilización de herramientas de "
              "inteligencia artificial», al final")


def hueco(clave: str) -> str:
    """Lo que rellena sellar_uso_ia. Legible tal cual, por si nadie lo sella."""
    return "⟦IA:" + clave + " · pendiente: lo completa la aplicación al archivar el papel⟧"


# Lo que el equipo completa: etiqueta y alto de fila (en Excel)
VERIFICACION = (("Qué se ha contrastado", 30), ("Con qué documentación", 30),
                ("Alcance de la verificación", 30),
                ("Diferencias encontradas y tratamiento", 30))
VERIFICADO = (("Nombre", 18), ("Fecha", 18))
EFECTO = (("Efecto", 45),)
REVISION = (("Revisado por", 18), ("Fecha", 18),
            ("¿Consta la evidencia y la verificación es efectiva?", 18),
            ("Observaciones / puntos a revisar", 45))
SI_NO = "¿Consta la evidencia y la verificación es efectiva?"

# Tipos de fila: 'dato' (lo deja la aplicacion), 'gris' (nota u orientacion, no es un dato),
# 'linea' (continuacion de la fila anterior), 'equipo' (en blanco, lo completa el equipo)


def contenido(*, papel: str, procedimiento: str, descripcion: str, alcance: str,
              resultado: str, comprobaciones: list, propuesta: str, generado: str,
              informacion_adicional: list | None = None) -> list:
    """Las ocho secciones del bloque: [(titulo, quien, [(tipo, etiqueta, valor, alto)])].

    papel          cabecera del papel: cliente, fecha y referencia
    procedimiento  el skill, como se llama en el plugin
    descripcion    qué hace el procedimiento, en una frase
    alcance        qué se pidió en esta ejecución (ejercicios, cuentas, prueba…)
    resultado      qué contiene el papel, y qué partes ha redactado la herramienta
    comprobaciones lo que comprobó la propia herramienta: va como resultado, no como verificación
    propuesta      qué verificar, orientativo; el alcance lo decide el equipo
    generado       fecha de generación
    informacion_adicional  lo facilitado que el MCP no ve (un PDF leído, p. ej.)
    """
    del papel  # va en la cabecera del bloque, no en una seccion
    # los skills reciben --generado en dos formatos; el bloque lo dice siempre igual (20/10/2026:
    # «2026-10-20» en unos papeles y «20/10/2026» en otros)
    m = re.fullmatch(r"\s*(\d{4})-(\d{2})-(\d{2})\s*", str(generado))
    if m:
        generado = m.group(3) + "/" + m.group(2) + "/" + m.group(1)
    s3 = [("dato", "Información facilitada", hueco("informacion"), 60)]
    s3 += [("linea", "", "• " + x, None) for x in (informacion_adicional or [])]
    s3 += [("dato", "Entorno de ejecución", hueco("entorno"), 30),
           ("dato", "Registro", "Cada exportación consta en AsistenteIA\\RegistroEmision.jsonl del "
                                "expediente (columnas y filas, nunca valores).", None)]
    s4 = [("dato", "Resultado", resultado, None),
          ("gris", "Comprobaciones de la propia herramienta",
           "No sustituyen la verificación del equipo (apartado 5).", None)]
    s4 += [("linea", "", "· " + c, None) for c in comprobaciones]
    s5 = [("gris", "Propuesta del procedimiento", "Orientativa; el alcance lo decide el equipo. "
           + propuesta, max(30, 15 * (1 + len(propuesta) // 110)))]
    s5 += [("equipo", e, "", a) for e, a in VERIFICACION]
    return [
        ("1. HERRAMIENTA", "Lo deja la aplicación", [
            ("dato", "Aplicación", hueco("aplicacion"), None),
            ("dato", "Modelo", hueco("modelo"), None),
            ("dato", "Plugin", hueco("plugin"), None),
            ("dato", "Procedimiento automatizado", procedimiento, None),
            ("dato", "Ejecutado por", hueco("usuario"), None),
            ("dato", "Fecha de generación", generado, None)]),
        ("2. PROCEDIMIENTO Y ALCANCE PEDIDO", "Lo deja la aplicación", [
            ("dato", "Procedimiento", descripcion, None),
            ("dato", "Alcance", alcance, None)]),
        ("3. INFORMACIÓN FACILITADA Y TRATAMIENTO", "Lo deja la aplicación", s3),
        ("4. RESULTADO", "Lo deja la aplicación", s4),
        ("5. VERIFICACIÓN DEL EQUIPO", "La completa el equipo", s5),
        ("6. VERIFICADO POR", "Lo completa el equipo", [("equipo", e, "", a) for e, a in VERIFICADO]),
        ("7. EFECTO EN LA CONCLUSIÓN DEL PAPEL", "Lo completa el equipo",
         [("equipo", e, "", a) for e, a in EFECTO]),
        ("8. REVISIÓN", "Supervisión: jefe de equipo o auditor responsable",
         [("equipo", e, "", a) for e, a in REVISION]),
    ]


# ── Excel ─────────────────────────────────────────────────────────────────────

def hoja_uso_ia(wb, *, papel: str, **campos):
    """Añade la hoja «Uso de IA» como PRIMERA del libro y la devuelve. Campos: contenido()."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    negrita, titulo_f = Font(bold=True), Font(bold=True, size=12)
    gris = Font(italic=True, color="595959")
    aviso = PatternFill("solid", fgColor="F8CBAD")
    blanco = PatternFill("solid", fgColor="FFF2CC")
    seccion_f = PatternFill("solid", fgColor="D9E1F2")
    fino = Side(style="thin", color="BFBFBF")
    caja = Border(left=fino, right=fino, top=fino, bottom=fino)
    envolver = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet(HOJA, 0)
    wb.active = 0
    for col, ancho in (("A", 34), ("B", 44), ("C", 30), ("D", 40)):
        ws.column_dimensions[col].width = ancho

    def celda(f, c, v, font=None, fill=None):
        x = ws.cell(f, c, v)
        x.alignment = envolver
        if font:
            x.font = font
        if fill:
            x.fill = fill
        return x

    celda(1, 1, TITULO, titulo_f)
    celda(2, 1, papel)
    for c in range(1, 5):
        ws.cell(4, c).fill = aviso
    celda(4, 1, PENDIENTE, negrita, aviso)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=4)
    ws.row_dimensions[4].height = 32

    f = 5
    si_no = None
    for titulo, quien, filas in contenido(papel=papel, **campos):
        f += 1
        for c in range(1, 5):
            ws.cell(f, c).fill = seccion_f
        celda(f, 1, titulo, negrita)
        celda(f, 4, quien, gris)
        f += 1
        for tipo, etiqueta, valor, alto in filas:
            celda(f, 1, etiqueta, gris if tipo == "gris" else negrita)
            if tipo == "equipo":
                for c in range(2, 5):
                    ws.cell(f, c).fill = blanco
                    ws.cell(f, c).border = caja
                if etiqueta == SI_NO:
                    si_no = f
            else:
                celda(f, 2, valor, gris if tipo == "gris" else None)
            ws.merge_cells(start_row=f, start_column=2, end_row=f, end_column=4)
            if alto:
                ws.row_dimensions[f].height = alto
            f += 1
    if si_no:
        dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("B" + str(si_no))
    ws.sheet_view.showGridLines = False
    return ws


def marcar_papel(ws, celda_etiqueta: str, celda_valor: str) -> None:
    """La marca en la hoja del papel, junto a Realizado y Verificado: que se vea al abrirlo."""
    from openpyxl.styles import Font, PatternFill

    ws[celda_etiqueta] = "Uso de IA"
    ws[celda_etiqueta].font = Font(bold=True)
    ws[celda_valor] = MARCA_PAPEL
    ws[celda_valor].fill = PatternFill("solid", fgColor="F8CBAD")


# ── Word ──────────────────────────────────────────────────────────────────────

def _sombrear(celda, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(shd)


def apartado_uso_ia(doc, *, papel: str, **campos) -> None:
    """Añade al final del documento el apartado del bloque, con un salto de página delante.

    Cada valor va en UNA sola celda y un solo run: sellar_uso_ia reemplaza el hueco dentro
    de word/document.xml, y un hueco partido entre runs no se encontraria."""
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, RGBColor

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(TITULO, level=1)
    p = doc.add_paragraph()
    r = p.add_run(papel)
    r.font.size = Pt(9)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = PENDIENTE
    c.paragraphs[0].runs[0].bold = True
    _sombrear(c, "F8CBAD")

    for titulo, quien, filas in contenido(papel=papel, **campos):
        h = doc.add_heading(titulo, level=2)
        q = h.add_run("   · " + quien)
        q.italic = True
        q.font.size = Pt(9)
        q.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for tipo, etiqueta, valor, _ in filas:
            fila = t.add_row().cells
            fila[0].text = etiqueta
            if fila[0].paragraphs[0].runs:
                fila[0].paragraphs[0].runs[0].bold = tipo != "gris"
                fila[0].paragraphs[0].runs[0].italic = tipo == "gris"
            if tipo == "equipo":
                # vacia de verdad: ni una pista dentro, que se leeria como respondida
                _sombrear(fila[1], "FFF2CC")
            else:
                fila[1].text = valor
                if tipo == "gris" and fila[1].paragraphs[0].runs:
                    run = fila[1].paragraphs[0].runs[0]
                    run.italic = True
                    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            for celda in fila:
                for par in celda.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9)
        for fila in t.rows:
            fila.cells[0].width = Pt(150)
            fila.cells[1].width = Pt(320)
