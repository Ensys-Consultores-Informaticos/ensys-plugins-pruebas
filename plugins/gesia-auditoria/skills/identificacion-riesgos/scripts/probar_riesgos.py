"""Arnes del catalogo de riesgos. No necesita Gesia, ni red, ni un master en disco.

    python probar_riesgos.py
    python probar_riesgos.py --catalogo "<trabajo>/riesgos.json"

Existe por un fallo concreto, del 04/09/2026: los dos riesgos de inclusion
obligatoria estaban fijados por numero -OBLIGATORIOS = (18, 27)-, y ese numero es
la numeracion interna de CADA master. Contra el master 25 acertaba; contra el
master 21 metia en el informe "Perdida de la concesion administrativa" y
"Obsolescencia del producto en catalogo" rotulados como obligatorios por la
NIA-ES 240, y dejaba fuera los dos que si lo son. Sin un solo aviso.

Los casos sinteticos reproducen las dos numeraciones medidas ese dia. Con
--catalogo se comprueba ademas un catalogo de verdad, el que deje
extraer_catalogo.py.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib_riesgos import (           # noqa: E402
    CLAVES_OBLIGATORIOS,
    SinObligatorios,
    anio_master,
    aviso_master,
    cargar_catalogo,
    ids_obligatorios,
    nota_master,
    resolver_obligatorios,
    salida_utf8,
)

VENTAS = "Integridad de las ventas (fraude)"
ELUSION = "Elusión de controles por la dirección (por sesgo o fraude)"


def catalogo(version, pares, relleno=()):
    """Un catalogo minimo: {id: nombre} y de que master dice venir."""
    riesgos = [{"id": i, "nombre": n, "area": "G", "area_nombre": ""}
               for i, n in list(pares) + list(relleno)]
    return {"version": version, "origen": version + ".gs3", "riesgos": riesgos}


# Las dos numeraciones reales, medidas el 04/09/2026.
M25 = catalogo("Master 25 CON RIESGOS", [(18, VENTAS), (27, ELUSION)])
M21 = catalogo("Master 21 CON RIESGOS", [(97, VENTAS), (28, ELUSION)],
               # Y estos son los que ocupaban el 18 y el 27 en ese master: los
               # que el codigo viejo colaba como obligatorios.
               [(18, "Pérdida de la concesión administrativa"),
                (27, "Obsolescencia del producto en catálogo (de carácter general)")])


def casos():
    yield ("el master 25 resuelve 18 y 27",
           lambda: ids_obligatorios(M25) == (18, 27))
    yield ("el master 21 resuelve 97 y 28",
           lambda: ids_obligatorios(M21) == (97, 28))
    yield ("los nombres resueltos son los dos correctos, en los dos masters",
           lambda: all([r["nombre"] for r in resolver_obligatorios(m)] == [VENTAS, ELUSION]
                       for m in (M25, M21)))

    # LA REGRESION. Si alguien vuelve a fijar (18, 27), esto muerde.
    yield ("contra el master 21 los obligatorios NO son 18 y 27",
           lambda: ids_obligatorios(M21) != (18, 27))
    yield ("y el 18 y el 27 existen en el master 21 siendo otra cosa: por eso "
           "fallaba en silencio",
           lambda: {r["id"] for r in M21["riesgos"]} >= {18, 27}
                   and 18 not in ids_obligatorios(M21))

    # El nombre se busca por un trozo, asi que un parentesis distinto no rompe.
    yield ("un nombre con otro parentesis sigue resolviendo",
           lambda: ids_obligatorios(catalogo(
               "Master X", [(3, "Integridad de las ventas"),
                            (9, "Elusión de controles por la dirección")])) == (3, 9))

    # Y cuando NO se puede resolver, se aborta. Nunca se sigue con lo que haya.
    yield ("un catalogo sin la elusion aborta",
           lambda: aborta(catalogo("Master Y", [(1, VENTAS)])))
    yield ("un catalogo con la elusion repetida aborta",
           lambda: aborta(catalogo("Master Z", [(1, VENTAS), (2, ELUSION), (3, ELUSION)])))
    yield ("un expediente renumerado, que no trae el catalogo, aborta",
           lambda: aborta(catalogo("Expediente", [(5, VENTAS)],
                                   [(1, "Relevancia Tesorería"), (2, "Relevancia compras")])))
    yield ("el mensaje de aborto dice de que catalogo habla",
           lambda: "Master Y" in mensaje(catalogo("Master Y", [(1, VENTAS)])))

    # El año del master, que es lo que avisa de haber cogido el de otro ejercicio.
    yield ("el año sale del nombre del master",
           lambda: (anio_master(M25), anio_master(M21)) == (2025, 2021))
    yield ("sin año en el nombre no se inventa",
           lambda: anio_master({"version": "CON RIESGOS", "origen": ""}) is None)
    yield ("un master de 2021 en un encargo de 2025 avisa",
           lambda: "2021" in (aviso_master(M21, 2025) or ""))
    yield ("el master del año que toca no avisa",
           lambda: aviso_master(M25, 2025) is None)
    yield ("sin ejercicio no se avisa de nada",
           lambda: aviso_master(M21, None) is None)
    # El master POSTERIOR al ejercicio es lo normal: la firma mantiene solo el vigente
    # (registro del 21/10/2026). Linea INFO, ni aviso, ni salida 1, ni informe.
    yield ("el master de 2025 en un encargo de 2024 no avisa: es lo habitual",
           lambda: aviso_master(M25, 2024) is None
           and "es lo habitual" in (nota_master(M25, 2024) or ""))
    yield ("y el del año que toca ni siquiera da nota",
           lambda: nota_master(M25, 2025) is None and nota_master(M21, 2025) is None)
    yield ("validar con el master 2025 y el ejercicio 2024: salida 0, INFO y sin avisos",
           lambda: master_vigente() == (0, True, []))

    # Las claves se comparan sin tildes, asi que ellas mismas no pueden llevarlas.
    yield ("las claves de busqueda van sin tildes y en minuscula",
           lambda: all(c == c.lower() and c.isascii() for c in CLAVES_OBLIGATORIOS))

    # El apartado de uso de IA (docs/evidencia-uso-ia.md), en Word
    yield ("el apartado de uso de IA lleva los huecos de la aplicacion, cada uno en un solo run",
           lambda: huecos_docx() == ["aplicacion", "entorno", "informacion", "modelo", "plugin", "usuario"])
    yield ("y los 11 apartados del equipo salen vacios: ni una pista dentro",
           lambda: celdas_equipo() == (11, 0))
    yield ("lib_uso_ia.py es identico en todos los skills que lo llevan",
           lambda: copias_iguales())

    # Cada script arranca con `python -I` desde otra carpeta: en Cowork se lanzan asi y
    # sin la ruta del script en sys.path no encuentran lib_riesgos (20/10/2026).
    for script in SCRIPTS:
        yield ("python -I " + script + " --help arranca desde otra carpeta",
               lambda s=script: lanzar(s, "--help")[0] == 0)

    # El nombre es opcional: sin el se usa el del catalogo y NO es un aviso (20/10/2026:
    # contaba como aviso y una seleccion correcta salia con codigo 1).
    yield ("una seleccion sin nombres valida con salida 0, una linea INFO y sin avisos",
           lambda: sin_nombre() == (0, True, []))

    # Los avisos de la validacion llegan al informe, por fichero y no a mano.
    yield ("un riesgo sin evidencia da salida 1 y el aviso queda en el JSON",
           lambda: (lambda r: r[0] == 1 and len(r[1]) == 1
                    and "evidencia cuantitativa" in r[1][0])(con_aviso()))
    yield ("y el aviso consta en el informe Word",
           lambda: aviso_en_informe())
    yield ("generar_informe no acepta la validacion de otra seleccion",
           lambda: informe_con_otra_validacion() == 2)
    yield ("ni una validacion fallida",
           lambda: informe_con_validacion_fallida() == 2)
    yield ("y sin --avisos no arranca",
           lambda: lanzar("generar_informe.py", "--analisis", "x", "--seleccion", "x",
                          "--generado", "x", "--ref", "x", "--salida", "x")[0] == 2)

    # La referencia del papel sale de Referencias de ESTE expediente (registro del
    # 21/10/2026: un codigo propuesto de memoria era alli el de otro papel).
    skill_md = (AQUI.parent / "SKILL.md").read_text(encoding="utf-8")
    yield ("el SKILL.md hace la segunda busqueda amplia de la referencia, con % de JET",
           lambda: "LIKE '%riesgo%'" in skill_md and "LIKE '%planificaci%n%'" in skill_md)
    yield ("y comprueba en Referencias el codigo que de el usuario",
           lambda: "WHERE CodigoReferencia = '<el código que ha dado>'" in skill_md)
    yield ("y no trae ningun codigo de referencia de ejemplo que se pueda copiar",
           lambda: not re.search(r"\b[A-Z]{1,3}\)\d", skill_md))
    yield ("con --ref vacia el informe sale y la celda REF. P.T. queda en blanco",
           lambda: ref_vacia() == (0, ""))

    # El directorio de trabajo en Cowork: las variables no pasan de una llamada a otra y
    # un rm -rf con la ruta calculada lo bloquea la seguridad (registro del 21/10/2026).
    yield ("el SKILL.md imprime la ruta del directorio de trabajo al crearlo",
           lambda: 'mktemp -d "${TMPDIR:-/tmp}/riesgos-XXXX")" && echo "$TRABAJO"' in skill_md)
    yield ("y lo borra con la ruta literal, nunca con la variable ni con $(cat …)",
           lambda: 'rm -rf "<la ruta que anotaste>"' in skill_md
           and 'rm -rf "$TRABAJO"' not in skill_md and 'rm -rf "$(cat' not in skill_md)

    # Registro del 21/10/2026: la narrativa se redacto abriendo los JSON y leyendo
    # mal los importes con coma decimal. La tabla va al resumen, ya convertida.
    yield ("analisis_resumen.md trae la tabla de cifras para la narrativa, ya en numero",
           lambda: tabla_en_resumen())
    yield ("y el SKILL.md manda redactar la narrativa con ella, sin abrir los JSON",
           lambda: "Cifras para la narrativa" in skill_md
           and "No abras `analisis.json`" in skill_md)
    # El rotulo de la clase de entidad dice de donde viene, no que lo decidio el auditor
    yield ("--mercantil se rotula «indicado con --mercantil», sin atribuirlo al auditor",
           lambda: (lambda r: r[0] == 0 and "(indicado con --mercantil)" in r[1]
                    and "auditor" not in r[1])(analizar("--mercantil")))
    yield ("y --sin-animo-de-lucro igual",
           lambda: (lambda r: r[0] == 0 and "indicado con --sin-animo-de-lucro" in r[1]
                    and "auditor" not in r[1])(analizar("--sin-animo-de-lucro")))
    yield ("el SKILL.md pide confirmar la clase de entidad antes de pasar el flag",
           lambda: "se confirman con el auditor" in skill_md)
    # Papel que ya existe: se pregunta, nunca se sobrescribe sin preguntar
    yield ("el SKILL.md comprueba si el informe ya existe y ofrece el sufijo de rehecho",
           lambda: "(rehecho AAAA-MM-DD)" in skill_md and "device_list_dir" in skill_md
           and "`force`" in skill_md)
    # MCP 1.26.1: configurar() sugiere el ultimo master usado en el equipo
    yield ("el SKILL.md propone el último máster usado si configurar lo sugiere",
           lambda: "master_riesgos" in skill_md and "El último usado en este equipo" in skill_md)


def _doc_uso_ia():
    import docx
    from lib_uso_ia import apartado_uso_ia

    d = docx.Document()
    apartado_uso_ia(d, papel="CLIENTE · Auditoría a 31/12/25 · X)1", procedimiento="prueba",
                    descripcion="d", alcance="a", resultado="r", comprobaciones=["c1"],
                    propuesta="p", generado="2026-10-19")
    return d


def huecos_docx() -> list:
    import re

    d = _doc_uso_ia()
    runs = [r.text for t in d.tables for row in t.rows for c in row.cells
            for p in c.paragraphs for r in p.runs]
    return sorted(m.group(1) for t in runs for m in re.finditer(r"⟦IA:([a-z_]+)[^⟧]*⟧", t))


def celdas_equipo():
    from lib_uso_ia import EFECTO, REVISION, VERIFICACION, VERIFICADO

    etiquetas = {e for e, _ in VERIFICACION + VERIFICADO + EFECTO + REVISION}
    filas = [r for t in _doc_uso_ia().tables for r in t.rows if r.cells[0].text in etiquetas]
    return len(filas), sum(1 for r in filas if r.cells[1].text.strip())


def copias_iguales() -> bool:
    propia = Path(__file__).resolve().parent / "lib_uso_ia.py"
    otras = [p for p in propia.parents[2].glob("*/scripts/lib_uso_ia.py") if p.resolve() != propia]
    return bool(otras) and all(p.read_bytes() == propia.read_bytes() for p in otras)


AQUI = Path(__file__).resolve().parent
SCRIPTS = ("analizar_cifras.py", "analizar_ra.py", "extraer_catalogo.py", "generar_informe.py",
           "mostrar_riesgos.py", "validar_seleccion.py", "probar_riesgos.py")
# riesgos que no son obligatorios, para llegar al minimo orientativo (5) y que el
# aviso de «pocos riesgos» no se mezcle con lo que se mide
CUENTAS = "Relevancia de cuentas a cobrar (Clientes y Deudores)"
M25_COMPLETO = catalogo("Master 25 CON RIESGOS", [(18, VENTAS), (27, ELUSION)],
                        [(33, CUENTAS), (40, "Relevancia de existencias"),
                         (41, "Relevancia de tesorería")])


def _con_cifra(i: int) -> dict:
    return {"id": i, "justificacion": "La partida varía.",
            "evidencia": {"epigrafe": "A.B", "concepto": "Partida", "importe": 1.0}}


CON_CIFRA = [_con_cifra(40), _con_cifra(41)]


def lanzar(script: str, *args) -> tuple:
    """Lanza un script con `python -I` desde una carpeta que no es la suya."""
    with tempfile.TemporaryDirectory() as fuera:
        entorno = dict(os.environ, PYTHONIOENCODING="utf-8")
        r = subprocess.run([sys.executable, "-I", str(AQUI / script), *args], cwd=fuera,
                           capture_output=True, text=True, encoding="utf-8", env=entorno)
    return r.returncode, r.stdout + r.stderr


def _validar(dir_: Path, riesgos: list, cat=M25_COMPLETO, ejercicio="2025") -> tuple:
    """Valida `riesgos` contra `cat`; devuelve (salida, stdout, ruta del JSON de avisos)."""
    (dir_ / "riesgos.json").write_text(json.dumps(cat, ensure_ascii=False), encoding="utf-8")
    (dir_ / "seleccion.json").write_text(json.dumps({"riesgos": riesgos}, ensure_ascii=False),
                                         encoding="utf-8")
    avisos = dir_ / "avisos.json"
    cod, out = lanzar("validar_seleccion.py", "--seleccion", str(dir_ / "seleccion.json"),
                      "--catalogo", str(dir_ / "riesgos.json"), "--ejercicio", ejercicio,
                      "--avisos-salida", str(avisos))
    return cod, out, avisos


def master_vigente() -> tuple:
    with tempfile.TemporaryDirectory() as d:
        cod, out, avisos = _validar(Path(d), OBLIGATORIOS_OK + CON_CIFRA + [_con_cifra(33)],
                                    ejercicio="2024")
        datos = json.loads(avisos.read_text(encoding="utf-8"))
    linea = [x for x in out.splitlines() if x.startswith("INFO") and "máster vigente (2025)" in x]
    return cod, bool(linea), datos["avisos"]


# Cifras sinteticas con los importes como llegan del API: texto con coma decimal.
# Saldo1..5 = ya con el signo de Gesia (entidades del MCP).
def _fila(clase, seccion, epi, concepto_, s1, s2, campo="Concepto"):
    f = {"CodigoClase": clase, "CodigoSeccion": seccion, "CodigoEpigrafe": epi,
         "CodigoDesglose": "", campo: concepto_}
    for i, v in enumerate((s1, s2, s2, s2, s2), 1):
        f["Saldo" + str(i)] = v
    return f


BALANCE = [
    _fila("A", "TG", "", "TOTAL ACTIVO", "1000000,00", "800000,00"),
    _fila("A", "B", "", "B) ACTIVO CORRIENTE", "600000,00", "500000,00"),
    _fila("A", "B", "3", "Deudores comerciales y otras cuentas a cobrar", "125432,10", "100000,00"),
    _fila("A", "B", "7", "Efectivo y otros activos líquidos equivalentes", "5000,00", "4000,00"),
    _fila("P", "TG", "", "TOTAL PATRIMONIO NETO Y PASIVO", "1000000,00", "800000,00"),
]
PYG = [_fila(*x, campo="ConceptoIngresos") for x in (
    ("I", "1", "1", "1. Importe neto de la cifra de negocios", "2000000,00", "1500000,00"),
    ("I", "1", "4", "4. Aprovisionamientos", "1200000,00", "900000,00"),
    ("ST", "2", "", "A.1) RESULTADO DE EXPLOTACIÓN", "-50000,00", "30000,00"),
    ("ST", "4", "", "A.4) RESULTADO DEL EJERCICIO PROCEDENTE DE OPERACIONES CONTINUADAS",
     "-40000,00", "20000,00"),
    ("T", "5", "", "A.5) RESULTADO DEL EJERCICIO (A.4+19)", "-40000,00", "20000,00"),
)]


def analizar(*flags) -> tuple:
    """analizar_cifras sobre las cifras sinteticas: (salida, stdout, resumen)."""
    with tempfile.TemporaryDirectory() as d:
        dir_ = Path(d)
        (dir_ / "b.json").write_text(json.dumps(BALANCE, ensure_ascii=False), encoding="utf-8")
        (dir_ / "p.json").write_text(json.dumps(PYG, ensure_ascii=False), encoding="utf-8")
        cod, out = lanzar("analizar_cifras.py", "--balance", str(dir_ / "b.json"),
                          "--pyg", str(dir_ / "p.json"), "--ir-t", "50000",
                          "--ir-t-origen", "gesia", "--ejercicios", "2024,2023,2022,2021,2020",
                          "--salida", str(dir_ / "analisis.json"), *flags)
        res = dir_ / "analisis_resumen.md"
        return cod, out, res.read_text(encoding="utf-8") if res.exists() else ""


def tabla_en_resumen() -> bool:
    cod, _out, res = analizar()
    tabla = res.split("## Cifras para la narrativa (2024 frente a 2023)", 1)[-1]
    return (cod == 0 and tabla != res
            and "| 125 432 | 100 000 | +25 432 | +25,4 % |" in tabla      # coma decimal leida
            and "Importe neto de la cifra de negocios | 2 000 000 | 1 500 000" in tabla
            and "+33,3 %" in tabla
            and "Efectivo" not in tabla                                    # bajo la IR_T
            and "RESULTADO DEL EJERCICIO (A.4+19) | -40 000 | 20 000" in tabla
            and "PROCEDENTE" not in tabla
            and "cambia de signo" in tabla)


def _informe(dir_: Path, avisos: Path) -> tuple:
    (dir_ / "analisis.json").write_text("{}", encoding="utf-8")
    salida = dir_ / "informe.docx"
    cod, out = lanzar("generar_informe.py", "--analisis", str(dir_ / "analisis.json"),
                      "--seleccion", str(dir_ / "seleccion.json"),
                      "--catalogo", str(dir_ / "riesgos.json"), "--avisos", str(avisos),
                      "--cliente", "CLIENTE SINTETICO", "--cierre", "31/12/25",
                      "--generado", "2026-10-20", "--ref", "X)1", "--salida", str(salida))
    return cod, salida


OBLIGATORIOS_OK = [{"id": 18, "justificacion": "NIA-ES 240."},
                   {"id": 27, "justificacion": "NIA-ES 240."}]


def sin_nombre() -> tuple:
    with tempfile.TemporaryDirectory() as d:
        cod, out, avisos = _validar(Path(d), OBLIGATORIOS_OK + CON_CIFRA + [_con_cifra(33)])
        datos = json.loads(avisos.read_text(encoding="utf-8"))
    return cod, "INFO" in out and "sin nombre" in out, datos["avisos"]


def con_aviso() -> tuple:
    with tempfile.TemporaryDirectory() as d:
        cod, _out, avisos = _validar(Path(d), OBLIGATORIOS_OK + CON_CIFRA + [
            {"id": 33, "justificacion": "Sin cifra."}])
        datos = json.loads(avisos.read_text(encoding="utf-8"))
    return cod, datos["avisos"]


def aviso_en_informe() -> bool:
    import docx

    with tempfile.TemporaryDirectory() as d:
        dir_ = Path(d)
        _cod, _out, avisos = _validar(dir_, OBLIGATORIOS_OK + CON_CIFRA
                                      + [{"id": 33, "justificacion": "x"}])
        cod, salida = _informe(dir_, avisos)
        if cod != 0:
            return False
        texto = "\n".join(p.text for p in docx.Document(str(salida)).paragraphs)
    return "evidencia cuantitativa" in texto and "Avisos de la validación" in texto


def ref_vacia() -> tuple:
    """Informe con --ref "": (salida, texto de la celda REF. P.T.)."""
    import docx

    with tempfile.TemporaryDirectory() as d:
        dir_ = Path(d)
        _cod, _out, avisos = _validar(dir_, OBLIGATORIOS_OK + CON_CIFRA + [_con_cifra(33)])
        (dir_ / "analisis.json").write_text("{}", encoding="utf-8")
        salida = dir_ / "informe.docx"
        cod, _ = lanzar("generar_informe.py", "--analisis", str(dir_ / "analisis.json"),
                        "--seleccion", str(dir_ / "seleccion.json"),
                        "--catalogo", str(dir_ / "riesgos.json"), "--avisos", str(avisos),
                        "--cliente", "CLIENTE SINTETICO", "--cierre", "31/12/25",
                        "--generado", "2026-10-21", "--ref", "", "--salida", str(salida))
        if cod != 0:
            return cod, None
        celdas = [r.cells[1].text for r in docx.Document(str(salida)).tables[0].rows
                  if r.cells[0].text == "REF. P.T."]
    return cod, celdas[0] if celdas else None


def informe_con_otra_validacion() -> int:
    with tempfile.TemporaryDirectory() as d:
        dir_ = Path(d)
        _validar(dir_, OBLIGATORIOS_OK + [{"id": 33, "justificacion": "x"}])
        # la seleccion se corrige DESPUES de validar: el JSON de avisos es de la otra
        (dir_ / "seleccion.json").write_text(json.dumps({"riesgos": OBLIGATORIOS_OK}),
                                             encoding="utf-8")
        return _informe(dir_, dir_ / "avisos.json")[0]


def informe_con_validacion_fallida() -> int:
    with tempfile.TemporaryDirectory() as d:
        dir_ = Path(d)
        # falta la elusion de controles: la validacion falla (salida 2)
        _validar(dir_, [{"id": 18, "justificacion": "x"}])
        return _informe(dir_, dir_ / "avisos.json")[0]


def aborta(cat) -> bool:
    try:
        resolver_obligatorios(cat)
    except SinObligatorios:
        return True
    return False


def mensaje(cat) -> str:
    try:
        resolver_obligatorios(cat)
    except SinObligatorios as exc:
        return str(exc)
    return ""


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    p.add_argument("--catalogo", help="un riesgos.json de verdad, para comprobarlo también")
    args = p.parse_args()

    fallos = 0
    for nombre, prueba in casos():
        try:
            bien = bool(prueba())
        except Exception as exc:                       # noqa: BLE001
            bien, nombre = False, nombre + "  [" + type(exc).__name__ + ": " + str(exc)[:80] + "]"
        print(("  ok   " if bien else "  FALLA ") + nombre)
        fallos += 0 if bien else 1

    if args.catalogo:
        print("\nCatálogo real: " + args.catalogo)
        try:
            cat = cargar_catalogo(args.catalogo)
            obl = resolver_obligatorios(cat)
            print("  máster: " + str(cat.get("version") or "sin identificar")
                  + "  ·  " + str(len(cat["riesgos"])) + " riesgos"
                  + "  ·  año " + str(anio_master(cat) or "?"))
            for r in obl:
                print("  ok   obligatorio " + str(r["id"]) + " — " + r["nombre"])
        except Exception as exc:                       # noqa: BLE001
            print("  FALLA " + type(exc).__name__ + ": " + str(exc)[:200])
            fallos += 1

    print("\n" + ("Todo bien." if not fallos else str(fallos) + " comprobación(es) fallan."))
    return 0 if not fallos else 1


if __name__ == "__main__":
    sys.exit(main())
