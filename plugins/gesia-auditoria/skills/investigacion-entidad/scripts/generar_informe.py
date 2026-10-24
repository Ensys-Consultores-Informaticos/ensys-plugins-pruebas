# -*- coding: utf-8 -*-
"""Informe de investigación de la entidad en Word.

    python generar_informe.py --contenido informe.json \
        --evidencia-societario "....json" --evidencia-sector-publico "....json" \
        --evidencia-riesgos "....json" \
        --cliente "RAZON SOCIAL, S.A." --cif A00000000 \
        --cif-origen "BORME (enlace), confirmado por el auditor" \
        --registro "Registro Mercantil (BORME)" --cierre 31/12/24 \
        --ventana "31/12/2019 - 26/08/2026" --generado 2026-08-26 --ref "" \
        --salida "<expediente>/AsistenteIA/InvestigacionEntidad/....docx"

`--generado` es obligatorio: nada en este proyecto lee el reloj.

La cabecera lleva los mismos campos y en el mismo orden que los papeles de
continuidad de saldos e identificación de riesgos, para que los papeles del
expediente se lean igual.

La sección de fuentes NO la redacta el modelo: se construye aquí uniendo los
`fuentes_consultadas` de los JSON que devolvieron los tres agentes. Así el informe no
puede citar una fuente que los agentes no registraran.

Esos JSON son entrada de trabajo, no se archivan: la evidencia es el propio papel, con
cada hecho citado a su fuente y cada fuente con su fecha de consulta en la sección 8.
En Cowork los agentes no escriben ficheros y su JSON llega como mensaje, que el modelo
copia al directorio de trabajo; con `--comprobar` se coteja cada copia con el
`recuento` que declaró el agente, antes de redactar nada:

    python generar_informe.py --comprobar \
        --evidencia-societario "....json" --evidencia-sector-publico "....json" \
        --evidencia-riesgos "....json"

Y termina con el apartado «Utilización de herramientas de inteligencia artificial»
(lib_uso_ia.py, docs/evidencia-uso-ia.md), con su marca en la cabecera: lo de la
herramienta lo escribe este script, con huecos que rellena sellar_uso_ia del MCP al
archivar; lo del equipo -verificación, quién, efecto y revisión- sale en blanco.
"""

from __future__ import annotations

import argparse
import json
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# Los modulos hermanos se buscan junto a este fichero: con `python -I` la carpeta
# del script no entra en sys.path y el import fallaba (registro del 20/10/2026).
sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib_uso_ia import MARCA_DOCX, apartado_uso_ia
from validar_cif import FORMAS, normalizar, validar

MARCAS = "✓⚠○"

# Lo que el equipo deberia contrastar, para el apartado de uso de IA. Describe, no dictamina:
# el alcance de la verificacion lo decide el equipo (docs/evidencia-uso-ia.md).
PROPUESTA = (
    "Abrir las fuentes citadas de los hechos relevantes para el encargo y contrastar que dicen lo "
    "que el informe recoge; confirmar que la entidad que aparece en cada fuente es la auditada, por "
    "denominación y CIF, y no un homónimo; y valorar qué información falta por fuentes no indexadas "
    "o de pago y si conviene obtenerla por otra vía. La evaluación del efecto en el encargo es del "
    "auditor."
)

# Las fuentes publicas a las que viajan la denominacion y el CIF, para el apartado 3 del bloque.
# Lo que el MCP no ve -las busquedas las hacen los agentes, no el servidor- lo declara el papel.
FUENTES_BUSQUEDA = "PLACE, BDNS, CENDOJ, BOE y boletines oficiales, y prensa"

# Limitaciones ESTRUCTURALES de las fuentes públicas gratuitas: son de las
# fuentes, no de la entidad, y se repiten en todos los informes. Texto fijo para
# que dos clientes distintos reciban la misma redacción; el modelo solo marca
# cuáles aplican por su id. Verificadas en la calibración con expediente real
# del 26/08/2026.
CATALOGO_LIMITACIONES = {
    "borme_sin_historico": (
        "No existe un buscador gratuito que indexe el histórico completo de actos "
        "del BORME por entidad; los actos citados proceden de búsquedas puntuales "
        "sobre boe.es y pueden no cubrir todos los publicados. Pueden existir "
        "actos intermedios no detectados, incluidos depósitos de cuentas."),
    "registro_mercantil_pago": (
        "El detalle del Registro Mercantil (situación vigente de cargos, depósito "
        "de cuentas, capital social) es de pago; la parte gratuita es escasa. Los "
        "huecos societarios de este informe son atribuibles a esa limitación, no "
        "a la entidad."),
    "aeat_censal": (
        "El estado censal en la AEAT exige certificado electrónico y no es "
        "consultable por terceros; no se ha verificado."),
    "place_sin_api": (
        "PLACE (contrataciondelestado.es) no ofrece un servicio público de "
        "consulta directa utilizable por herramientas automatizadas; la "
        "comprobación se hace por búsqueda web restringida al dominio, con un "
        "margen de incertidumbre residual mayor que el de BDNS."),
    "cendoj_no_fetchable": (
        "El buscador de CENDOJ no devuelve resultados legibles de forma fiable a "
        "herramientas automatizadas; la conclusión «sin datos» se apoya en "
        "búsqueda combinada, no en una exploración manual exhaustiva del buscador "
        "nativo."),
    "boletines_no_fetchable": (
        "Los buscadores de boletines autonómicos y provinciales presentan la "
        "misma limitación de lectura automatizada que el CENDOJ."),
    "prensa_no_indexada": (
        "La ausencia de piezas en cabeceras de prensa nacional se basa en los "
        "resultados del motor de búsqueda utilizado; no puede descartarse que "
        "existan piezas no indexadas."),
    "cnae_no_oficial": (
        "El CNAE no se publica en ninguna fuente registral gratuita; cuando "
        "consta, procede de agregadores (fiabilidad media)."),
}

DECLARACION = (
    "La información de este documento la ha recopilado un modelo de lenguaje en "
    "fuentes públicas de internet, con cada hecho citado a su fuente y cada fuente "
    "consultada, con su fecha de consulta, en la sección 8. No es una "
    "conclusión de auditoría: la evaluación del efecto de esta información en el "
    "encargo es del auditor, que es quien firma.")

LEYENDA = ("✓ hecho verificado en la fuente citada · ⚠ indicio, resultado del "
           "análisis cruzado · ○ sin información: consta que se buscó y no hay "
           "dato público")

# Orden de presentación de la tabla de fuentes: primero las que dieron datos
# (las «consultadas con éxito»), después las vacías y al final las inaccesibles.
ORDEN_RESULTADO = ["con datos", "sin datos", "inaccesible"]

# Lo que cuenta como hecho en el JSON de cada agente, para cotejar la copia con el
# `recuento` que declara (agents/investigador-*.md): listas que cuentan por entradas, y
# objetos que cuentan uno si no son null.
HECHOS = {
    "Societario": (("administradores", "apoderados", "cambios_societarios"),
                   ("constitucion", "actividad_declarada", "cnae", "cuentas_depositadas")),
    "Sector público": (("contratos_publicos", "subvenciones"), ()),
    "Riesgos legales": (("procedimientos_judiciales", "sanciones", "prensa"), ()),
}

# Una busqueda en el motor no es una pagina: va con la url vacia y el nombre empezando asi
# (contrato de los agentes). En el registro del 21/10/2026 las filas «Busqueda web general:
# ... sentencia» llevaban la URL de una ficha que salio entre los resultados, y el papel la
# presentaba como fuente de la busqueda.
PREFIJO_BUSQUEDA = "busqueda web"
SIN_URL_BUSQUEDA = "— (búsqueda general, sin página concreta)"


def salida_utf8() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass


def parrafo(doc, texto: str, cursiva=False, tam=10, negrita=False):
    p = doc.add_paragraph()
    t = p.add_run(texto)
    t.italic, t.bold = cursiva, negrita
    t.font.size = Pt(tam)
    return p


def tabla(doc, cabeceras: list, filas: list, estilo="Light Grid Accent 1"):
    """Tabla con o sin fila de cabecera (sin cabecera si todas vienen vacías)."""
    con_cabecera = any(str(c).strip() for c in cabeceras)
    t = doc.add_table(rows=1 if con_cabecera else 0, cols=len(cabeceras))
    t.style = estilo
    if con_cabecera:
        for i, c in enumerate(cabeceras):
            celda = t.rows[0].cells[i]
            celda.text = str(c)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = "" if v is None else str(v)
    return t


def validar_contenido(contenido: dict) -> list:
    """Devuelve la lista de defectos. Con defectos NO se genera informe."""
    defectos = []
    for nombre, items in (contenido.get("secciones") or {}).items():
        for item in items:
            if item.get("marca") not in tuple(MARCAS):
                defectos.append("la sección '%s' lleva un item con marca '%s'; "
                                "las válidas son ✓ ⚠ ○"
                                % (nombre, item.get("marca")))
            if item.get("marca") == "✓" and not item.get("fuente"):
                defectos.append("hecho ✓ sin fuente en la sección '%s': «%s…»"
                                % (nombre, str(item.get("texto", ""))[:60]))
    for clave in contenido.get("limitaciones_catalogo", []):
        if clave not in CATALOGO_LIMITACIONES:
            defectos.append("id de limitación desconocido: '%s'. No se inventan "
                            "limitaciones de catálogo" % clave)
    if not contenido.get("limitaciones_catalogo") and \
            not contenido.get("limitaciones_especificas"):
        defectos.append("el apartado de limitaciones no puede quedar vacío")
    return defectos


def items_seccion(doc, items: list):
    """Escribe los items de una sección con su marca y su fuente."""
    if not items:
        parrafo(doc, "○ Sin contenido en esta sección.", tam=10)
        return
    for item in items:
        texto = item["marca"] + " " + item["texto"]
        if item.get("fuente"):
            texto += " (" + item["fuente"] + ")"
        if item.get("incertidumbre"):
            texto += " — incertidumbre " + item["incertidumbre"]
        parrafo(doc, texto, tam=10)


def cargar_evidencias(rutas: dict) -> tuple:
    """Lee los JSON de los agentes. Devuelve ({bloque: datos o None}, defectos). Un
    bloque sin ruta es un agente degradado (None); un JSON que no carga es un defecto:
    casi siempre, una copia del mensaje del agente cortada o retocada."""
    datos, defectos = {}, []
    for bloque, ruta in rutas.items():
        if not ruta:
            datos[bloque] = None
            continue
        try:
            # utf-8-sig: una copia escrita desde Windows puede llevar BOM, y no es un defecto
            d = json.loads(Path(ruta).read_text(encoding="utf-8-sig"))
        except FileNotFoundError:
            defectos.append("no existe la copia del JSON del agente de " + bloque.lower()
                            + ": " + str(ruta))
            continue
        except ValueError as exc:
            defectos.append("la copia del JSON del agente de " + bloque.lower() + " no carga ("
                            + str(exc) + "): vuelve a copiar su mensaje entero y sin retocar")
            continue
        if not isinstance(d, dict) or not isinstance(d.get("fuentes_consultadas", []), list):
            defectos.append("el JSON del agente de " + bloque.lower() + " no tiene la forma del "
                            "contrato (un objeto con la lista fuentes_consultadas)")
            continue
        datos[bloque] = d
    return datos, defectos


def contar_hechos(bloque: str, d: dict) -> int:
    listas, objetos = HECHOS.get(bloque, ((), ()))
    return (sum(len(d.get(k) or []) for k in listas)
            + sum(1 for k in objetos if d.get(k)))


def _entero(v):
    if isinstance(v, bool):
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def cotejar_recuento(datos: dict) -> tuple:
    """Coteja cada copia con el `recuento` que declaro su agente. Devuelve (defectos,
    lineas, sin_recuento): un descuadre es defecto -la copia ha perdido o ganado algo-;
    un agente que no declaro recuento no aborta, pero el papel lo dice."""
    defectos, lineas, sin_recuento = [], [], []
    for bloque, d in datos.items():
        if d is None:
            continue
        n_f, n_h = len(d.get("fuentes_consultadas") or []), contar_hechos(bloque, d)
        rec = d.get("recuento")
        if not isinstance(rec, dict):
            sin_recuento.append(bloque)
            lineas.append(bloque + ": JSON válido · " + str(n_f) + " fuentes · " + str(n_h)
                          + " hechos · SIN recuento declarado: compáralo a ojo con su mensaje")
            continue
        texto = bloque + ": JSON válido · " + str(n_f) + " fuentes"
        decl_f = _entero(rec.get("fuentes_consultadas"))
        if decl_f is None:
            defectos.append("el recuento del agente de " + bloque.lower() + " no trae el número "
                            "de fuentes_consultadas")
        elif decl_f != n_f:
            defectos.append("la copia del agente de " + bloque.lower() + " tiene " + str(n_f)
                            + " fuentes y el agente declaró " + str(decl_f) + ": vuelve a "
                            "copiar su mensaje entero y sin retocar")
        texto += " (declaradas " + str(rec.get("fuentes_consultadas")) + ") · " + str(n_h) + " hechos"
        if rec.get("hechos") is not None:
            decl_h = _entero(rec.get("hechos"))
            if decl_h != n_h:
                defectos.append("la copia del agente de " + bloque.lower() + " tiene " + str(n_h)
                                + " hechos y el agente declaró " + str(rec.get("hechos"))
                                + ": vuelve a copiar su mensaje entero y sin retocar")
            texto += " (declarados " + str(rec.get("hechos")) + ")"
        lineas.append(texto)
    return defectos, lineas, sin_recuento


def es_busqueda_general(nombre) -> bool:
    plano = "".join(c for c in unicodedata.normalize("NFD", str(nombre or ""))
                    if unicodedata.category(c) != "Mn")
    return plano.strip().lower().startswith(PREFIJO_BUSQUEDA)


def fuentes_de_evidencia(datos: dict) -> tuple:
    """Une los fuentes_consultadas de los agentes. Fila: bloque, nombre, url, fecha,
    resultado. Una evidencia ausente (agente degradado) se anota. Una busqueda general
    sale sin URL aunque el agente pusiera una. Devuelve (filas, urls_quitadas)."""
    filas, quitadas = [], 0
    for bloque, d in datos.items():
        if d is None:
            filas.append((bloque, "— agente degradado: sin evidencia —",
                          "", "", "inaccesible"))
            continue
        for f in d.get("fuentes_consultadas", []):
            nombre, url = f.get("nombre", "—"), str(f.get("url") or "").strip()
            if es_busqueda_general(nombre):
                quitadas += 1 if url else 0
                url = SIN_URL_BUSQUEDA
            elif not url:
                url = "—"
            filas.append((bloque, nombre, url,
                          f.get("fecha_consulta", ""), f.get("resultado", "—")))

    def clave_orden(fila):
        resultado = fila[4]
        for i, prefijo in enumerate(ORDEN_RESULTADO):
            if str(resultado).startswith(prefijo):
                return i
        return len(ORDEN_RESULTADO)

    return sorted(filas, key=clave_orden), quitadas


def avisos_de_evidencia(datos: dict) -> list:
    """Los avisos de cada agente, en orden: [(bloque, texto)]. Van al papel tal cual, en la
    seccion 8: ya no hay JSON archivado donde leerlos (registro del 21/10/2026)."""
    return [(bloque, str(a)) for bloque, d in datos.items() if d
            for a in (d.get("avisos") or [])]


def comprobaciones_uso_ia(args, secciones: dict, contenido: dict, evidencias: dict,
                          filas_fuentes: list, con_exito: int, sin_recuento=(),
                          urls_quitadas: int = 0) -> list:
    """Lo que la propia herramienta comprueba sobre ESTE informe: cada frase se calcula aquí
    sobre lo que el papel lleva, y si algo falla lo dice con «ATENCIÓN». Va en el resultado
    del apartado de uso de IA, nunca como verificación del equipo."""
    items = [i for lista in secciones.values() for i in lista]
    hechos = [i for i in items if i.get("marca") == "✓"]
    indicios = [i for i in items if i.get("marca") == "⚠"]
    vacios = [i for i in items if i.get("marca") == "○"]
    out = []

    # el CIF del informe: control, persona juridica y el registro que le toca por su letra
    cif = normalizar(args.cif)
    tipo, valido, motivo = validar(cif)
    if tipo == "cif" and valido:
        forma, registro, _ = FORMAS.get(cif[:1], (None, None, False))
        out.append("El CIF " + cif + " supera el carácter de control y es de persona jurídica ("
                   + str(forma) + ").")
        if registro == args.registro:
            out.append("El registro citado en el informe («" + args.registro + "») es el que "
                       "corresponde a la letra del CIF.")
        else:
            out.append("ATENCIÓN: el registro citado en el informe («" + args.registro + "») no es "
                       "el que corresponde a la letra del CIF («" + str(registro) + "»).")
    else:
        out.append("ATENCIÓN: el identificador " + cif + " no es un CIF válido de persona jurídica ("
                   + (motivo or "es de persona física") + ").")

    # el validador aborta con un hecho ✓ sin fuente o una marca invalida: si se llega aqui, es cierto
    out.append(("Hechos verificados (✓): " + str(len(hechos)) + ", todos con su fuente citada"
                if hechos else "No hay hechos verificados (✓)")
               + ". Todas las afirmaciones llevan una marca válida (✓ ⚠ ○); indicios (⚠): "
               + str(len(indicios)) + "; sin información (○): " + str(len(vacios)) + ".")
    sin_incert = [i for i in indicios if not i.get("incertidumbre")]
    if not indicios:
        out.append("No hay indicios (⚠) cuya incertidumbre declarar.")
    elif not sin_incert:
        out.append("Todos los indicios (⚠) declaran su incertidumbre.")
    else:
        out.append("ATENCIÓN: " + str(len(sin_incert)) + " de " + str(len(indicios))
                   + " indicios (⚠) no declaran su incertidumbre.")

    out.append("Limitaciones estructurales citadas: "
               + str(len(contenido.get("limitaciones_catalogo", [])))
               + ", todas textos fijos del procedimiento; propias de esta entidad: "
               + str(len(contenido.get("limitaciones_especificas", []))) + ".")

    # la seccion 8, desde lo que devolvieron los agentes: cuantas fuentes, y que agente no
    # dejo nada. Los JSON no se archivan (registro del 21/10/2026): lo que acreditan esta aqui.
    out.append("La sección 8 se ha construido desde las respuestas de los agentes, no desde la "
               "redacción. Fuentes: " + str(len(filas_fuentes)) + "; con datos: "
               + str(con_exito) + ".")
    cotejados = [b for b, d in evidencias.items() if d is not None and b not in sin_recuento]
    if cotejados:
        out.append("La respuesta de cada agente se ha cotejado con el recuento de fuentes y hechos "
                   "que declaró (" + ", ".join(b.lower() for b in cotejados) + "): cuadra.")
    for bloque in sin_recuento:
        out.append("ATENCIÓN: el agente de " + bloque.lower() + " no declaró recuento: su "
                   "respuesta no se ha podido cotejar.")
    if urls_quitadas:
        out.append("ATENCIÓN: " + str(urls_quitadas) + " búsquedas web generales traían la URL de "
                   "una página concreta; en la sección 8 constan sin URL.")
    for bloque, d in evidencias.items():
        if d is None:
            out.append("ATENCIÓN: sin evidencia del agente de " + bloque.lower() + ": su bloque "
                       "consta como degradado en la sección 8.")
        elif not any(f[0] == bloque for f in filas_fuentes):
            out.append("ATENCIÓN: la evidencia del agente de " + bloque.lower() + " no registra "
                       "fuentes consultadas.")
    avisos = avisos_de_evidencia(evidencias)
    if avisos:
        cuenta = {}
        for b, _ in avisos:
            cuenta[b] = cuenta.get(b, 0) + 1
        out.append("Avisos de los agentes: " + str(len(avisos)) + " ("
                   + ", ".join(b.lower() + " " + str(n) for b, n in cuenta.items())
                   + "), recogidos tal cual en la sección 8.")
    else:
        out.append("Los agentes no dejaron avisos.")
    return out


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    # Solo coteja las copias de los JSON de los agentes con su recuento, sin informe: va
    # justo despues de copiarlas, antes de redactar (registro del 21/10/2026).
    p.add_argument("--comprobar", action="store_true",
                   help="solo comprobar los JSON de los agentes; no genera informe")
    p.add_argument("--contenido",
                   help="JSON con las secciones escritas por el modelo")
    p.add_argument("--evidencia-societario")
    p.add_argument("--evidencia-sector-publico")
    p.add_argument("--evidencia-riesgos")
    p.add_argument("--cliente")
    p.add_argument("--cif")
    # De donde salio el CIF. La investigacion parte del nombre del expediente y el CIF lo busca
    # en internet y lo confirma el auditor: Gesia no lo guarda (la columna NIF de DatosAuditoria
    # es, desde hace varias versiones, el sector; 22/10/2026). Un informe sobre la entidad
    # equivocada es el peor resultado posible, asi que el papel dice como se identifico, y no
    # se puede omitir.
    p.add_argument("--cif-origen",
                   help="p. ej. 'indicado por el auditor' o "
                        "'BORME (enlace), confirmado por el auditor'")
    # El registro que corresponde a la entidad por la letra de su CIF (validar_cif.py).
    # La cabecera decia «BORME/BOE» a todo el mundo, tambien a una fundacion (05/10/2026).
    p.add_argument("--registro",
                   help="el campo 'registro' de validar_cif.py, p. ej. 'Registro Mercantil (BORME)'")
    p.add_argument("--cierre")
    p.add_argument("--ventana",
                   help="ventana temporal investigada, p. ej. '31/12/2019 - 26/08/2026'")
    p.add_argument("--generado")
    # La referencia sale de la tabla Referencias de ESTE expediente, y se pregunta
    # enseñando las candidatas; si no tiene una, se pasa vacía y la celda queda en
    # blanco. Nunca un codigo de memoria o de otro expediente: registro del 21/10/2026,
    # el codigo propuesto era alli el de otro papel.
    p.add_argument("--ref",
                   help="referencia del papel, comprobada en Referencias del expediente; "
                        "cadena vacía si el usuario no la tiene")
    p.add_argument("--salida")
    args = p.parse_args()

    rutas = {
        "Societario": args.evidencia_societario,
        "Sector público": args.evidencia_sector_publico,
        "Riesgos legales": args.evidencia_riesgos,
    }
    evidencias, defectos = cargar_evidencias(rutas)
    defectos_rec, lineas_rec, sin_recuento = cotejar_recuento(evidencias)
    defectos += defectos_rec

    if args.comprobar:
        if not any(rutas.values()):
            p.error("--comprobar necesita al menos un --evidencia-*")
        for linea in lineas_rec:
            print(linea)
        for d in defectos:
            print("DEFECTO: " + d)
        if defectos:
            return 2
        print("Copias comprobadas: " + ("cuadran con lo que declararon los agentes."
                                        if not sin_recuento else
                                        "sin descuadres; " + ", ".join(sin_recuento)
                                        + " sin recuento declarado."))
        return 0

    faltan = [o for o in ("contenido", "cliente", "cif", "cif_origen", "registro", "cierre",
                          "ventana", "generado", "ref", "salida") if getattr(args, o) is None]
    if faltan:
        p.error("faltan: " + ", ".join("--" + o.replace("_", "-") for o in faltan))

    contenido = json.loads(Path(args.contenido).read_text(encoding="utf-8"))
    defectos += validar_contenido(contenido)
    if defectos:
        for d in defectos:
            print("DEFECTO: " + d)
        return 2

    secciones = contenido.get("secciones") or {}

    doc = Document()
    doc.core_properties.title = "Investigación de la entidad en fuentes públicas"

    # ── Cabecera común de los papeles del expediente ──────────────────────────
    parrafo(doc, args.cliente, negrita=True, tam=12)
    parrafo(doc, "AUDITORIA A " + args.cierre)
    h = doc.add_heading("INVESTIGACIÓN DE LA ENTIDAD EN FUENTES PÚBLICAS", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    parrafo(doc, "Fuente: fuentes públicas de internet (" + args.registro + ", BOE y boletines oficiales, PLACE, BDNS, "
                 "CENDOJ y prensa); cada hecho cita su fuente, y la sección 8 recoge las "
                 "fuentes consultadas con su fecha de consulta")
    # Realizado y Verificado van en blanco: los rellena el auditor.
    tabla(doc, ["", ""], [["REF. P.T.", args.ref],
                          ["Realizado", ""],
                          ["Verificado", ""],
                          ["Uso de IA", MARCA_DOCX]])
    parrafo(doc, "Conocimiento de la entidad y su entorno por fuentes externas · "
                 "NIA-ES 315", cursiva=True, tam=9)
    parrafo(doc, "Generado: " + args.generado + " · Ventana investigada: "
                 + args.ventana, tam=9)
    parrafo(doc, LEYENDA, cursiva=True, tam=9)
    parrafo(doc, DECLARACION, cursiva=True, tam=9)

    # ── 1. Identificación ─────────────────────────────────────────────────────
    doc.add_heading("1. Identificación de la entidad", level=2)
    filas = [["Denominación", args.cliente], ["CIF", args.cif],
             ["Origen del CIF", args.cif_origen],
             ["Cierre del ejercicio auditado", args.cierre]]
    filas += [[i.get("campo", "—"), i.get("valor", "—")]
              for i in contenido.get("identificacion", [])]
    tabla(doc, ["", ""], filas)

    # ── 2 a 5: los bloques de hechos ──────────────────────────────────────────
    doc.add_heading("2. Estructura societaria", level=2)
    items_seccion(doc, secciones.get("societaria", []))
    doc.add_heading("3. Actividad y contexto sectorial", level=2)
    items_seccion(doc, secciones.get("actividad", []))
    doc.add_heading("4. Relación con el sector público", level=2)
    items_seccion(doc, secciones.get("sector_publico", []))
    doc.add_heading("5. Riesgos legales y reputacionales", level=2)
    items_seccion(doc, secciones.get("riesgos_legales", []))

    # ── 6. Análisis ───────────────────────────────────────────────────────────
    doc.add_heading("6. Análisis — elementos a contrastar", level=2)
    analisis = secciones.get("analisis", [])
    items_seccion(doc, analisis)
    if analisis:
        parrafo(doc, "Los elementos anteriores cruzan hechos ya citados; no son "
                     "conclusiones. Incertidumbre baja = respaldado por dos o más "
                     "fuentes; alta = hipótesis a partir de una coincidencia.",
                cursiva=True, tam=9)

    # ── 7. Limitaciones (obligatorio, y no puede quedar vacío) ────────────────
    doc.add_heading("7. Limitaciones del análisis", level=2)
    for clave in contenido.get("limitaciones_catalogo", []):
        parrafo(doc, "· " + CATALOGO_LIMITACIONES[clave], tam=10)
    for texto in contenido.get("limitaciones_especificas", []):
        parrafo(doc, "· " + texto, tam=10)

    # ── 8. Fuentes consultadas, desde la evidencia y no desde el modelo ──────
    doc.add_heading("8. Fuentes consultadas", level=2)
    filas_fuentes, urls_quitadas = fuentes_de_evidencia(evidencias)
    con_exito = sum(1 for f in filas_fuentes if str(f[4]).startswith("con datos"))
    parrafo(doc, str(len(filas_fuentes)) + " fuentes consultadas, "
                 + str(con_exito) + " con datos. Primero las que aportaron "
                 "información, después las consultadas sin resultado y las "
                 "inaccesibles: todas acreditan dónde se buscó.", tam=10)
    tabla(doc, ["Bloque", "Fuente", "URL", "Fecha consulta", "Resultado"],
          [list(f) for f in filas_fuentes])
    # Los avisos de los agentes describen los limites de su busqueda; sin JSON archivado,
    # constan aqui tal cual los devolvieron.
    avisos = avisos_de_evidencia(evidencias)
    if avisos:
        parrafo(doc, "Avisos de los agentes sobre su búsqueda, tal como los devolvieron:", tam=10)
        for bloque, texto in avisos:
            parrafo(doc, "· " + bloque + ": " + texto, tam=9)

    # ── Evidencia del uso de IA (lib_uso_ia, docs/evidencia-uso-ia.md) ────────
    # Lo central aqui no pasa por el MCP: la denominacion y el CIF viajan a internet por los
    # agentes. Eso lo declara el papel en informacion_adicional; lo servido por el MCP
    # (contexto_expediente, las consultas de Referencias) lo lista sellar_uso_ia.
    degradados = [b.lower() for b, d in evidencias.items() if d is None]
    adicional = [
        "La denominación y el CIF de la entidad se han enviado a búsquedas en fuentes públicas de "
        "internet (" + args.registro + ", " + FUENTES_BUSQUEDA + ") a través de los agentes de "
        "búsqueda de la herramienta, que no tienen acceso al expediente; con ellos, la forma "
        "jurídica, la fecha de cierre y la ventana temporal. Lo que vuelve son datos públicos: "
        "los hechos que usa el informe constan en él con su fuente, y las fuentes consultadas, "
        "con su fecha, en la sección 8."]
    if "expediente" not in args.cif_origen.lower():
        adicional.append("El CIF no se tomó del expediente: para identificarlo se buscó la "
                         "denominación en fichas públicas de empresa. Origen del CIF: "
                         + args.cif_origen + ".")
    apartado_uso_ia(
        doc,
        papel=args.cliente + " · Auditoría a " + args.cierre + " · "
              + (args.ref + " " if args.ref else "")
              + "Investigación de la entidad en fuentes públicas",
        procedimiento="investigacion-entidad",
        descripcion=("Investigación de la entidad auditada en fuentes públicas de internet "
                     "—registro que le corresponde, contratación pública, subvenciones, "
                     "jurisprudencia, boletines oficiales y prensa— mediante tres agentes de "
                     "búsqueda, con cada hecho citado a su fuente y las fuentes consultadas, con "
                     "su fecha de consulta, en la sección 8."),
        alcance=("Entidad " + args.cliente + ", CIF " + args.cif + ". Registro: " + args.registro
                 + ". Ventana investigada: " + args.ventana + ". Bloques: societario, sector "
                 "público y riesgos legales"
                 + (" (sin evidencia: " + ", ".join(degradados) + ")" if degradados else "")
                 + "."),
        resultado=("Este documento: " + str(sum(len(v) for v in secciones.values()))
                   + " afirmaciones en las secciones 2 a 6, y " + str(len(filas_fuentes))
                   + " fuentes consultadas. El informe lo ha compuesto la herramienta a partir de "
                     "lo que devolvieron los agentes, con cada hecho citado a su fuente: los datos "
                     "añadidos a la identificación, la redacción de las secciones 2 a 6 —incluido "
                     "el análisis cruzado— y las limitaciones propias de la entidad son suyos. "
                     "La denominación y el CIF constan en la sección 1 con el origen del CIF. "
                     "Las limitaciones estructurales "
                     "de las fuentes son textos fijos del procedimiento, y la sección 8 se "
                     "construye desde las respuestas de los agentes."),
        comprobaciones=comprobaciones_uso_ia(args, secciones, contenido, evidencias,
                                             filas_fuentes, con_exito, sin_recuento,
                                             urls_quitadas),
        propuesta=PROPUESTA,
        generado=args.generado,
        informacion_adicional=adicional,
    )

    destino = Path(args.salida)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    print("Informe escrito en " + str(destino))
    print("  secciones: " + ", ".join(sorted(secciones)))
    print("  fuentes: " + str(len(filas_fuentes)) + " (" + str(con_exito)
          + " con datos)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
