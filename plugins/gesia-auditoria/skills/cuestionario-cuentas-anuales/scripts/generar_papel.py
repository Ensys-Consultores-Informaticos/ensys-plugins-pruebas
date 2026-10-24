"""Papel de trabajo en Word: que se respondio, por que y con que evidencia.

El .xlsx no tiene sitio para justificar nada — la columna `Comentario` trae la
referencia normativa del master y sobrescribirla destruye dato. Sin este
documento el criterio no consta en ninguna parte, y un papel de trabajo cuyo
criterio no consta es indefendible en una inspeccion.

    python generar_papel.py --cuestionario trabajo/cuestionario.json \
        --respuestas trabajo/respuestas.json --generado 2026-08-23 \
        --cliente "RAZON SOCIAL, S.A." --cierre "31/12/25" \
        --xlsx "AG20-01_CLIENTE_2025.xlsx" --pdf "Cuentas anuales 2025.pdf" \
        --salida papel.docx

`--generado` es obligatorio a proposito: nada en este proyecto lee el reloj. Un
papel de trabajo tiene que poder regenerarse identico dentro de dos años.

Termina con el apartado «Utilizacion de herramientas de inteligencia artificial»
(lib_uso_ia, docs/evidencia-uso-ia.md). Va SOLO aqui: el .xlsx de importacion no se
toca, porque lo lee el importador de Gesia y una hoja de mas podria romperlo. Con
`--xlsx` el papel coteja ese Excel contra sus propias respuestas y dice el resultado
entre sus comprobaciones; `--pdf` es el nombre del documento que leyo el modelo, que
el MCP no ve y por eso se declara aqui.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# Los modulos hermanos se buscan junto a este fichero: con `python -I` la carpeta
# del script no entra en sys.path y el import fallaba (registro del 20/10/2026).
sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib_cuestionario import (
    COLUMNAS,
    RESPUESTAS,
    entero_o_none,
    es_cabecera,
    leer_cuestionario,
    partir_en_secciones,
    salida_utf8,
    valor_celda,
)
from lib_uso_ia import MARCA_DOCX, apartado_uso_ia

# Cuantas preguntas se detallan por seccion en la tabla larga. Con 137 preguntas
# el documento ya son ~20 paginas; el tope evita que un cuestionario de 481
# (memoria consolidada) lo convierta en algo que nadie abre.
MAX_POR_SECCION = 40

# Cuantos puntos o diferencias se citan en una comprobacion que falla. Lo justo para
# localizarlos; el resto se cuenta.
MAX_CITA = 8

# La hoja del Excel de importacion, la misma que escribe generar_xlsx.py
HOJA_XLSX = "Guia"

# Los codigos como se escriben en el apartado de uso de IA, con su tilde
RESPUESTAS_LEGIBLES = {1: "Sí", 2: "No", 3: "No aplica", 4: "Pendiente"}

# Que significa «p. N» en las evidencias. La numeracion impresa de unas cuentas anuales
# no suele coincidir con la del fichero -portada, indice y hojas firmadas sin numero-, y
# sin fijar una el auditor no sabe a que pagina ir. Se fija la del PDF porque es la unica
# que el modelo ve siempre, tambien en un escaneo sin texto (prueba en Cowork, 20/10/2026).
NUMERACION_PAGINAS = ("Las del fichero PDF, contando su primera página como la 1; no la "
                      "numeración impresa en las cuentas")

# Lo que el equipo deberia contrastar. Describe, no dictamina: el alcance es suyo.
PROPUESTA = (
    "Contrastar contra la memoria las respuestas «Sí» y «No», comprobando que el desglose "
    "figura, o falta, donde dicen el motivo y la evidencia citada, y una selección de las «No "
    "aplica»; resolver las «Pendiente» con el conocimiento del encargo; y revisar el Excel de "
    "importación antes de importarlo en Gesia. La respuesta de cada pregunta es del auditor."
)

DECLARACION = (
    "Las respuestas de este cuestionario las ha propuesto un modelo de lenguaje a partir "
    "de las cuentas anuales aportadas. NO son una conclusion de auditoria: son una "
    "propuesta que el auditor revisa, corrige y asume antes de importarla al expediente. "
    "Ante cualquier duda sobre si un desglose figura o no en la memoria, la respuesta "
    "emitida es 4 (Pendiente), nunca 1 (Si)."
)

LIMITACIONES = [
    "Solo se ha revisado si el desglose figura en la memoria. No se juzga si su contenido "
    "es correcto, ni si las cifras que contiene cuadran con la contabilidad.",
    "No se han contrastado las respuestas contra los saldos del expediente. Una pregunta "
    "respondida «no aplica» sobre un epigrafe en el que el cliente tiene saldo no se "
    "detecta en esta version.",
    "El balance y la cuenta de perdidas y ganancias no se han cuadrado contra el PDF.",
    "Las respuestas dependen de lo que el documento aportado permita leer. Un desglose "
    "presente pero redactado de forma que no se localice se responde 4 (Pendiente).",
]


def parrafo(doc, texto: str, cursiva: bool = False, tam: int = 10):
    p = doc.add_paragraph()
    t = p.add_run(texto)
    t.italic = cursiva
    t.font.size = Pt(tam)
    return p


def _citar(puntos: list) -> str:
    """Los primeros puntos de una lista, y cuantos mas hay."""
    texto = ", ".join(str(x) for x in puntos[:MAX_CITA])
    return texto + (" y " + str(len(puntos) - MAX_CITA) + " más" if len(puntos) > MAX_CITA else "")


def _normalizar(valor):
    """Igual que probar_fidelidad.py: un retorno de carro escapado es la misma celda."""
    if isinstance(valor, str):
        return valor.replace("_x000D_", "\r").replace("\r\n", "\n").replace("\r", "\n")
    return valor


def cotejar_xlsx(ruta: str, filas: list, codigos: dict) -> list:
    """Diferencias entre el Excel de importacion y lo que dice este papel. Vacia si coincide.

    Comprueba lo que el importador de Gesia necesita que sea cierto: una sola hoja, las doce
    columnas del export en su orden, las filas del cuestionario, las once columnas copiadas
    del expediente sin cambios y `Respuesta` igual a la respuesta de este papel en cada
    pregunta (las cabeceras de seccion conservan la suya, que es vacia)."""
    from openpyxl import load_workbook

    try:
        libro = load_workbook(ruta)
    except Exception as e:  # cualquier fallo al abrirlo es un hallazgo, no un error del papel
        return ["no se ha podido abrir (" + type(e).__name__ + ")"]
    if libro.sheetnames != [HOJA_XLSX]:
        return ["las hojas son " + ", ".join(libro.sheetnames) + " y debería ser solo «"
                + HOJA_XLSX + "»"]
    hoja = libro[HOJA_XLSX]
    cabecera = [c.value for c in hoja[1]]
    if cabecera != COLUMNAS:
        return ["la fila de títulos no es la del export de Gesia"]
    if hoja.max_row - 1 != len(filas):
        return [str(hoja.max_row - 1) + " filas y el cuestionario tiene " + str(len(filas))]
    difs = []
    for i, fila in enumerate(filas, start=2):
        pregunta = not es_cabecera(fila)
        for j, columna in enumerate(COLUMNAS, start=1):
            if columna == "Respuesta" and pregunta:
                esperado = codigos.get(str(fila["CodigoOrden"]))
            else:
                esperado = valor_celda(columna, fila.get(columna))
            leido = hoja.cell(row=i, column=j).value
            if _normalizar(leido) != _normalizar(esperado) or type(leido) is not type(esperado):
                difs.append(columna + " de la fila " + str(i))
    return difs


def comprobaciones(cuestionario: dict, preguntas: list, lista: list, xlsx: str | None) -> list:
    """Lo que la herramienta comprueba sobre lo que lleva este papel. Cada frase es verdad
    calculada aqui mismo; si algo falla, lo dice con ATENCIÓN. Van en el resultado del bloque
    de uso de IA, nunca como verificacion: esa es del equipo."""
    def punto(fila):
        return str(fila["Punto"]) + "." + str(fila["DesglosePunto"])

    por_orden_preg = {str(f["CodigoOrden"]): f for f in preguntas}
    ordenes = [str(r.get("orden")) for r in lista]
    codigos = {str(r.get("orden")): entero_o_none(r.get("respuesta")) for r in lista}
    sin = [punto(f) for o, f in por_orden_preg.items() if o not in codigos]
    ajenas = sorted({o for o in ordenes if o not in por_orden_preg})
    repetidas = sorted({o for o in ordenes if ordenes.count(o) > 1})
    fuera = [o for o in ordenes if codigos.get(o) not in RESPUESTAS]

    salida = []
    n = len(preguntas)
    if not (sin or ajenas or repetidas):
        salida.append("Las " + str(n) + " preguntas del cuestionario tienen una respuesta, y "
                      "ninguna respuesta queda fuera del cuestionario.")
    else:
        partes = []
        if sin:
            partes.append("preguntas sin respuesta: " + str(len(sin)) + " (puntos " + _citar(sin) + ")")
        if ajenas:
            partes.append("respuestas que no corresponden a ninguna pregunta: " + str(len(ajenas))
                          + " (órdenes " + _citar(ajenas) + ")")
        if repetidas:
            partes.append("preguntas con más de una respuesta: " + str(len(repetidas))
                          + " (órdenes " + _citar(repetidas) + ")")
        salida.append("ATENCIÓN: el cuestionario tiene " + str(n) + " preguntas, y hay "
                      + "; ".join(partes) + ".")
    if not fuera:
        salida.append("Todas las respuestas usan los códigos de Gesia: 1 Sí, 2 No, 3 No aplica "
                      "o 4 Pendiente.")
    else:
        salida.append("ATENCIÓN: hay respuestas que no usan los códigos 1 a 4 de Gesia: "
                      + str(len(fuera)) + " (órdenes " + _citar(fuera) + ").")

    def falta(campo, codigo):
        return [punto(por_orden_preg[str(r["orden"])]) for r in lista
                if str(r.get("orden")) in por_orden_preg
                and entero_o_none(r.get("respuesta")) == codigo
                and not str(r.get(campo) or "").strip()]

    n_si = sum(1 for o in por_orden_preg if codigos.get(o) == 1)
    sin_evidencia = falta("evidencia", 1)
    salida.append(("Las respuestas «Sí» (" + str(n_si) + ") citan dónde está el desglose en la "
                   "memoria.") if not sin_evidencia else
                  ("ATENCIÓN: hay respuestas «Sí» que no citan dónde está el desglose: "
                   + str(len(sin_evidencia)) + " de " + str(n_si) + " (puntos "
                   + _citar(sin_evidencia) + ")."))
    n_no = sum(1 for o in por_orden_preg if codigos.get(o) == 2)
    n_pte = sum(1 for o in por_orden_preg if codigos.get(o) == 4)
    sin_motivo = falta("motivo", 2) + falta("motivo", 4)
    salida.append(("Las respuestas «No» (" + str(n_no) + ") y «Pendiente» (" + str(n_pte)
                   + ") explican su motivo.") if not sin_motivo else
                  ("ATENCIÓN: hay respuestas «No» o «Pendiente» que no explican su motivo: "
                   + str(len(sin_motivo)) + " de " + str(n_no + n_pte) + " (puntos "
                   + _citar(sin_motivo) + ")."))

    if xlsx:
        nombre = Path(xlsx).name
        difs = cotejar_xlsx(xlsx, cuestionario["filas"], codigos)
        if not difs:
            salida.append("El Excel de importación «" + nombre + "» tiene una sola hoja, con las "
                          + str(len(COLUMNAS)) + " columnas del formato de exportación de Gesia y "
                          "las " + str(len(cuestionario["filas"])) + " filas del cuestionario: las "
                          "columnas del expediente, copiadas sin cambios, y la columna Respuesta "
                          "igual a la de este papel.")
        else:
            salida.append("ATENCIÓN: el Excel de importación «" + nombre + "» no coincide con este "
                          "papel: " + _citar(difs) + ".")
    else:
        salida.append("El Excel de importación no se ha cotejado con este papel: no se indicó "
                      "al generarlo.")
    return salida


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--cuestionario", required=True)
    p.add_argument("--respuestas", required=True)
    p.add_argument("--generado", required=True, help="Fecha de generacion, AAAA-MM-DD")
    p.add_argument("--salida", required=True)
    # El cliente y el cierre los pasa quien llama, con lo que le dio
    # contexto_expediente. Antes se leian del API de Gesia, y desde Cowork eso
    # no responde: el papel salia con la cabecera en blanco y sin avisar de
    # nada, porque datos_encargo se comia la excepcion. Visto en cliente el
    # 23/08/2026.
    p.add_argument("--cliente", default="", help="razon social, de contexto_expediente")
    p.add_argument("--cierre", default="", help="fecha de cierre, de contexto_expediente")
    p.add_argument("--xlsx", help="el Excel de importacion de generar_xlsx.py, para cotejarlo")
    # Obligatorio: el papel dice siempre que documento se leyo. Con varias copias de las
    # cuentas en el expediente, un papel que no lo dice no deja saber contra cual se
    # contesto (prueba en Cowork, 20/10/2026).
    p.add_argument("--pdf", required=True,
                   help="el PDF de las cuentas anuales leido, tal como lo confirmo el usuario")
    args = p.parse_args()
    if not args.pdf.strip():
        print("ERROR: --pdf vacío. El papel tiene que decir qué PDF de las cuentas anuales "
              "se ha leído.", file=sys.stderr)
        return 2

    cuestionario = leer_cuestionario(args.cuestionario)
    crudo = json.loads(Path(args.respuestas).read_text(encoding="utf-8"))
    lista = crudo["respuestas"] if isinstance(crudo, dict) else crudo
    por_orden = {str(r["orden"]): r for r in lista}

    cliente, cierre = args.cliente.strip(), args.cierre.strip()
    filas = cuestionario["filas"]
    secciones = partir_en_secciones(filas)
    preguntas = [f for f in filas if not es_cabecera(f)]

    doc = Document()
    doc.add_heading("Revisión del contenido de la memoria", level=0)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub.add_run(cliente or "(cliente no disponible)").bold = True

    ficha = doc.add_table(rows=0, cols=2)
    ficha.style = "Light Grid Accent 1"
    for etiqueta, valor in [
        ("Cierre del ejercicio", cierre or "—"),
        ("Cuestionario", cuestionario["guia"]),
        ("Preguntas", str(len(preguntas))),
        ("Cuentas anuales leídas", args.pdf.strip()),
        ("Páginas citadas", NUMERACION_PAGINAS),
        ("Generado el", args.generado),
        ("Uso de IA", MARCA_DOCX),
    ]:
        celda = ficha.add_row().cells
        celda[0].text = etiqueta
        celda[1].text = valor

    doc.add_heading("Cómo leer este documento", level=1)
    parrafo(doc, DECLARACION, cursiva=True)

    # Resumen.
    doc.add_heading("Resumen", level=1)
    reparto = {c: 0 for c in RESPUESTAS}
    for r in lista:
        codigo = entero_o_none(r.get("respuesta"))
        reparto[codigo] = reparto.get(codigo, 0) + 1
    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = "Light Grid Accent 1"
    for i, texto in enumerate(("Código", "Significado", "Preguntas")):
        tabla.rows[0].cells[i].text = texto
    for codigo in sorted(RESPUESTAS):
        fila = tabla.add_row().cells
        fila[0].text = str(codigo)
        fila[1].text = RESPUESTAS[codigo]
        fila[2].text = str(reparto.get(codigo, 0))

    # Lo que el auditor tiene que mirar primero.
    for codigo, titulo, entradilla in [
        (2, "Desgloses que no figuran en la memoria",
         "Preguntas respondidas «No». Son las que exigen decidir si falta información "
         "en la memoria o si la pregunta no era aplicable."),
        (4, "Pendientes de resolver",
         "Preguntas que el documento aportado no permite responder. Requieren "
         "intervención del auditor antes de dar el cuestionario por cerrado."),
    ]:
        seleccion = [r for r in lista if entero_o_none(r.get("respuesta")) == codigo]
        doc.add_heading(f"{titulo} ({len(seleccion)})", level=1)
        parrafo(doc, entradilla, cursiva=True, tam=9)
        if not seleccion:
            parrafo(doc, "Ninguna.")
            continue
        for r in seleccion:
            fila = next((f for f in preguntas if str(f["CodigoOrden"]) == str(r["orden"])), None)
            if fila is None:
                continue
            enc = doc.add_paragraph()
            enc.add_run(f"{fila['Punto']}.{fila['DesglosePunto']}  ").bold = True
            enc.add_run(" ".join(str(fila["Descripcion"]).split()))
            if r.get("motivo"):
                parrafo(doc, f"Motivo: {r['motivo']}", tam=9)
            if r.get("evidencia"):
                parrafo(doc, f"Evidencia: {r['evidencia']}", tam=9)

    # Detalle completo.
    doc.add_heading("Detalle por secciones", level=1)
    for seccion in secciones:
        if not seccion["preguntas"]:
            continue
        doc.add_heading(
            f"Punto {seccion['punto']} — {' '.join(str(seccion['titulo']).split())}", level=2
        )
        tabla = doc.add_table(rows=1, cols=4)
        tabla.style = "Light Grid Accent 1"
        for i, texto in enumerate(("#", "Pregunta", "Resp.", "Motivo y evidencia")):
            tabla.rows[0].cells[i].text = texto
        for fila in seccion["preguntas"][:MAX_POR_SECCION]:
            r = por_orden.get(str(fila["CodigoOrden"]), {})
            codigo = entero_o_none(r.get("respuesta"))
            celdas = tabla.add_row().cells
            celdas[0].text = f"{fila['Punto']}.{fila['DesglosePunto']}"
            celdas[1].text = " ".join(str(fila["Descripcion"]).split())
            celdas[2].text = f"{codigo} {RESPUESTAS[codigo]}" if codigo in RESPUESTAS else "—"
            justificacion = " ".join(x for x in (r.get("motivo"), r.get("evidencia")) if x)
            celdas[3].text = justificacion
        sobran = len(seccion["preguntas"]) - MAX_POR_SECCION
        if sobran > 0:
            parrafo(doc, f"({sobran} preguntas más de esta sección, en el .xlsx)", cursiva=True, tam=9)

    doc.add_heading("Alcance y limitaciones", level=1)
    for texto in LIMITACIONES:
        doc.add_paragraph(texto, style="List Bullet")

    # ── Evidencia del uso de IA (lib_uso_ia, docs/evidencia-uso-ia.md) ────────
    # Las comprobaciones se calculan aqui, sobre lo que lleva el papel: no afirma lo que
    # otro script dijo, afirma lo que este ha visto
    comprobado = comprobaciones(cuestionario, preguntas, lista, args.xlsx)
    reparto_txt = " · ".join(f"{c} {RESPUESTAS_LEGIBLES[c]}: {reparto.get(c, 0)}"
                             for c in sorted(RESPUESTAS))
    xlsx_txt = "«" + Path(args.xlsx).name + "»" if args.xlsx else "de este mismo cuestionario"
    # el nombre, sin carpeta, tambien si la ruta es de Windows y el script corre en Linux
    pdf_txt = "«" + args.pdf.strip().replace("\\", "/").rsplit("/", 1)[-1] + "»"
    contestadas = cuestionario["n_contestadas"]
    apartado_uso_ia(
        doc,
        papel=(cliente or "(cliente no disponible)") + " · Auditoría a " + (cierre or "—") + " · "
              + cuestionario["guia"] + " Revisión del contenido de la memoria",
        procedimiento="cuestionario-cuentas-anuales",
        descripcion=("Revisión del contenido de la memoria: para cada pregunta del cuestionario de "
                     "cuentas anuales del expediente, la herramienta lee las cuentas anuales del "
                     "cliente y propone si el desglose figura en la memoria (1 Sí, 2 No, 3 No "
                     "aplica o 4 Pendiente), con el motivo y dónde lo ha visto."),
        alcance=("Cuestionario " + cuestionario["guia"] + " del expediente: "
                 + str(len(preguntas)) + " preguntas en "
                 + str(sum(1 for s in secciones if s["preguntas"])) + " secciones"
                 + (" (contestadas ya en el expediente antes de esta ejecución: "
                    + str(contestadas) + ")" if contestadas else "")
                 + ". Cuentas anuales " + pdf_txt + (", cierre " + cierre if cierre else "") + "."),
        resultado=("Este papel justifica las " + str(len(lista)) + " respuestas propuestas ("
                   + reparto_txt + "). Las respuestas, sus motivos y las referencias a la memoria "
                   "los ha redactado la herramienta; las preguntas son las del cuestionario del "
                   "expediente. Las mismas respuestas van en el Excel de importación " + xlsx_txt
                   + ", con el formato de exportación de Gesia, que el auditor revisa e importa "
                     "desde Gesia: este procedimiento no escribe en el expediente."),
        comprobaciones=comprobado,
        propuesta=PROPUESTA,
        generado=args.generado,
        informacion_adicional=[
            "Cuentas anuales del cliente en PDF (" + pdf_txt + ") — leídas directamente por el "
            "modelo, sin pasar por el servidor de Gesia — documento completo, en claro (sin "
            "seudonimizar)"],
    )

    salida = Path(args.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(salida)

    print(f"Escrito: {salida}")
    print("Reparto: " + " · ".join(
        f"{c} {RESPUESTAS[c]}: {reparto.get(c, 0)}" for c in sorted(RESPUESTAS)))
    faltan = len(preguntas) - len(lista)
    if faltan:
        print(f"AVISO: {faltan} preguntas sin entrada en respuestas.json")
    atencion = [c for c in comprobado if c.startswith("ATENCIÓN")]
    print(f"Uso de IA: {len(comprobado)} comprobaciones de la herramienta, "
          f"{len(atencion)} con ATENCIÓN" + (":" if atencion else "."))
    for c in atencion:
        print("  " + c)
    return 0


if __name__ == "__main__":
    sys.exit(main())
