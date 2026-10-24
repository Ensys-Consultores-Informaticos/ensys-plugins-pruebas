"""Arnes del papel de trabajo y del Excel de importacion. No necesita Gesia, ni red, ni PDF.

    python probar_papel.py
    python probar_papel.py --papel "<expediente>/AsistenteIA/CuestionarioCuentasAnuales/<papel>.docx"

Genera con un cuestionario sintetico -todas las columnas con valor, que un cuestionario
en blanco no prueba los tipos- el Excel y el papel en una carpeta temporal, y comprueba
el apartado «Utilizacion de herramientas de inteligencia artificial» (docs/evidencia-uso-ia.md):
que existe, que lleva exactamente los huecos que rellena sellar_uso_ia, cada uno en un solo
run, que los once apartados del equipo salen vacios, que sus comprobaciones son verdad
-y dicen ATENCION cuando no lo son- y que el Excel de importacion sigue siendo el export de
Gesia, sin hojas de mas. Y que lib_uso_ia.py es identico en todos los skills que lo llevan.

Con --papel comprueba ademas el apartado de un papel ya generado.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))

from lib_cuestionario import salida_utf8                   # noqa: E402
from lib_uso_ia import EFECTO, REVISION, TITULO, VERIFICACION, VERIFICADO  # noqa: E402

HUECOS = ["aplicacion", "entorno", "informacion", "modelo", "plugin", "usuario"]


def _fila(orden, punto, desglose, texto, respuesta=""):
    return {"Prioritaria": "True", "CodigoArea": "12", "CodigoGuia": "PRUEBA)01",
            "CodigoOrden": str(orden), "Punto": str(punto), "DesglosePunto": str(desglose),
            "Descripcion": texto, "CodigoClaseReferencia": "3", "CodigoReferencia": "PRUEBA)01",
            "CodigoRealizado": "1", "Respuesta": respuesta,
            "Comentario": "Art. " + str(orden) + "\r\nsegunda linea"}


CUESTIONARIO = [
    _fila(1, 1, "", "ACTIVIDAD DE LA EMPRESA"),
    _fila(2, 1, 1, "Se informa del domicilio y la forma legal", "1"),
    _fila(3, 1, 2, "Se describe la actividad principal"),
    _fila(4, 2, "", "BASES DE PRESENTACION"),
    _fila(5, 2, 1, "Se indica el marco normativo aplicado"),
    _fila(6, 2, 2, "Se explican las estimaciones significativas"),
    _fila(7, 2, 3, "Se informa de los cambios de criterio"),
]

BUENAS = [
    {"orden": 2, "respuesta": 1, "motivo": "Consta en la nota 1", "evidencia": "Nota 1, p. 5"},
    {"orden": 3, "respuesta": 2, "motivo": "La nota 1 no describe la actividad (saldo 1.000 EUR)"},
    {"orden": 5, "respuesta": 1, "motivo": "Nota 2.1", "evidencia": "Nota 2.1, p. 6"},
    {"orden": 6, "respuesta": 4, "motivo": "Depende de juicios que la memoria no detalla"},
    {"orden": 7, "respuesta": 3, "motivo": "No hay cambios de criterio"},
]


PDF_PRUEBA = "C:\\Expediente\\Cuentas\\Cuentas anuales de prueba.pdf"


def _generar(carpeta: Path, respuestas: list, respuestas_xlsx: list | None = None,
             con_xlsx: bool = True, pdf: str | None = PDF_PRUEBA) -> tuple:
    """Escribe el Excel y el papel. Devuelve (salidas de los procesos, ruta xlsx, ruta docx)."""
    (carpeta / "cuestionario.json").write_text(json.dumps(CUESTIONARIO), encoding="utf-8")
    (carpeta / "respuestas.json").write_text(json.dumps({"respuestas": respuestas}), encoding="utf-8")
    (carpeta / "resp_xlsx.json").write_text(
        json.dumps({"respuestas": respuestas_xlsx or respuestas}), encoding="utf-8")
    xlsx, docx_ = carpeta / "AG20-01_PRUEBA.xlsx", carpeta / "papel.docx"
    entorno = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    a = subprocess.run([sys.executable, str(AQUI / "generar_xlsx.py"),
                        "--cuestionario", str(carpeta / "cuestionario.json"),
                        "--respuestas", str(carpeta / "resp_xlsx.json"), "--salida", str(xlsx)],
                       capture_output=True, text=True, encoding="utf-8", env=entorno)
    orden = [sys.executable, str(AQUI / "generar_papel.py"),
             "--cuestionario", str(carpeta / "cuestionario.json"),
             "--respuestas", str(carpeta / "respuestas.json"), "--generado", "2026-10-19",
             "--cliente", "CLIENTE DE PRUEBA, S.A.", "--cierre", "31/12/25",
             "--salida", str(docx_)]
    if pdf is not None:
        orden += ["--pdf", pdf]
    if con_xlsx:
        orden += ["--xlsx", str(xlsx)]
    b = subprocess.run(orden, capture_output=True, text=True, encoding="utf-8", env=entorno)
    return (a, b), xlsx, docx_


def textos_runs(ruta) -> list:
    import docx

    d = docx.Document(str(ruta))
    return [r.text for t in d.tables for row in t.rows for c in row.cells
            for p in c.paragraphs for r in p.runs]


def tiene_apartado(ruta) -> bool:
    import docx

    return any(p.text == TITULO and p.style.name.startswith("Heading")
               for p in docx.Document(str(ruta)).paragraphs)


def huecos_docx(ruta) -> list:
    return sorted(m.group(1) for t in textos_runs(ruta) for m in re.finditer(r"⟦IA:([a-z_]+)[^⟧]*⟧", t))


def celdas_equipo(ruta) -> tuple:
    import docx

    etiquetas = {e for e, _ in VERIFICACION + VERIFICADO + EFECTO + REVISION}
    filas = [r for t in docx.Document(str(ruta)).tables for r in t.rows
             if len(r.cells) == 2 and r.cells[0].text in etiquetas]
    return len(filas), sum(1 for r in filas if r.cells[1].text.strip())


def _ficha(ruta) -> dict:
    """La ficha de cabecera: la primera tabla del papel, etiqueta -> valor."""
    import docx

    tabla = docx.Document(str(ruta)).tables[0]
    return {f.cells[0].text: f.cells[1].text for f in tabla.rows if len(f.cells) == 2}


def texto_tablas(ruta) -> str:
    import docx

    return "\n".join(c.text for t in docx.Document(str(ruta)).tables
                     for row in t.rows for c in row.cells)


def hojas(ruta) -> list:
    from openpyxl import load_workbook

    return load_workbook(str(ruta)).sheetnames


def copias_iguales() -> bool:
    propia = AQUI / "lib_uso_ia.py"
    otras = [p for p in propia.parents[2].glob("*/scripts/lib_uso_ia.py") if p.resolve() != propia]
    return bool(otras) and all(p.read_bytes() == propia.read_bytes() for p in otras)


def casos(tmp: Path):
    bien = tmp / "bien"
    bien.mkdir()
    (pa, pb), xlsx, papel = _generar(bien, BUENAS)
    yield ("los dos scripts terminan sin error", lambda: pa.returncode == 0 and pb.returncode == 0)
    yield ("el Excel de importacion sigue siendo el export de Gesia: una sola hoja, «Guia»",
           lambda: hojas(xlsx) == ["Guia"])
    yield ("y no lleva nada del bloque de uso de IA",
           lambda: _sin_texto_ia(xlsx))
    yield ("(a) el papel lleva el apartado de uso de IA", lambda: tiene_apartado(papel))
    yield ("(b) con los huecos de la aplicacion, cada uno en un solo run",
           lambda: huecos_docx(papel) == HUECOS)
    yield ("(c) y los 11 apartados del equipo salen vacios: ni una pista dentro",
           lambda: celdas_equipo(papel) == (11, 0))
    yield ("la marca de uso de IA va en la ficha de cabecera",
           lambda: "Pendiente de verificación" in texto_tablas(papel).split(TITULO)[0])
    yield ("con todo en orden, ninguna comprobacion dice ATENCIÓN",
           lambda: "ATENCIÓN" not in texto_tablas(papel))
    yield ("y el Excel se ha cotejado y coincide",
           lambda: "Excel de importación «AG20-01_PRUEBA.xlsx» tiene una sola hoja" in texto_tablas(papel))
    yield ("el resultado dice que existe el Excel y que la herramienta redacto las respuestas",
           lambda: all(x in texto_tablas(papel) for x in
                       ("Excel de importación «AG20-01_PRUEBA.xlsx»", "los ha redactado la herramienta")))
    yield ("el PDF leido consta como informacion facilitada en claro, por su nombre",
           lambda: "(«Cuentas anuales de prueba.pdf») — leídas directamente" in texto_tablas(papel)
           and "en claro (sin seudonimizar)" in texto_tablas(papel))
    yield ("la ficha dice que PDF se leyo, con la ruta tal como se indico",
           lambda: _ficha(papel).get("Cuentas anuales leídas") == PDF_PRUEBA)
    yield ("y en una linea, que paginas son las que citan las evidencias",
           lambda: _ficha(papel).get("Páginas citadas", "").startswith("Las del fichero PDF"))

    # Sin PDF no hay papel: tiene que decir siempre que documento se leyo
    for nombre, pdf in (("sin --pdf", None), ("con --pdf vacio", "  ")):
        carpeta = tmp / ("pdf_" + str(pdf is None))
        carpeta.mkdir()
        (_, sp), _, papel_sp = _generar(carpeta, BUENAS, pdf=pdf)
        yield (nombre + ", el papel no se genera (salida 2)",
               lambda sp=sp, papel_sp=papel_sp: sp.returncode == 2 and not papel_sp.exists())

    # Lo que tiene que hacer saltar las comprobaciones
    mal = tmp / "mal"
    mal.mkdir()
    malas = [dict(r) for r in BUENAS]
    malas[0].pop("evidencia")                     # un «Sí» sin evidencia
    malas[3]["motivo"] = ""                       # un «Pendiente» sin motivo
    otras = [dict(r) for r in BUENAS]
    otras[1]["respuesta"] = 3                     # el Excel dice otra cosa que el papel
    (_, mb), _, papel_mal = _generar(mal, malas[:-1], respuestas_xlsx=otras)
    yield ("un «Sí» sin evidencia sale con ATENCIÓN",
           lambda: "ATENCIÓN: hay respuestas «Sí» que no citan dónde está el desglose: 1 de 2 (puntos 1.1)" in texto_tablas(papel_mal))
    yield ("un «Pendiente» sin motivo sale con ATENCIÓN",
           lambda: "ATENCIÓN: hay respuestas «No» o «Pendiente» que no explican su motivo: 1 de 2 (puntos 2.2)" in texto_tablas(papel_mal))
    yield ("una pregunta sin respuesta sale con ATENCIÓN",
           lambda: "preguntas sin respuesta: 1 (puntos 2.3)" in texto_tablas(papel_mal))
    yield ("un Excel que no dice lo mismo que el papel sale con ATENCIÓN",
           lambda: "ATENCIÓN: el Excel de importación «AG20-01_PRUEBA.xlsx» no coincide"
           in texto_tablas(papel_mal))
    yield ("y el script lo dice en pantalla", lambda: "con ATENCIÓN:" in mb.stdout)

    sin = tmp / "sin"
    sin.mkdir()
    (_, _), _, papel_sin = _generar(sin, BUENAS, con_xlsx=False)
    yield ("sin --xlsx, el papel dice que el Excel no se ha cotejado",
           lambda: "no se ha cotejado con este papel" in texto_tablas(papel_sin))

    yield ("lib_uso_ia.py es identico en todos los skills que lo llevan", copias_iguales)


def _sin_texto_ia(xlsx) -> bool:
    from openpyxl import load_workbook

    libro = load_workbook(str(xlsx))
    return not any(isinstance(c.value, str) and ("Uso de IA" in c.value or "⟦IA:" in c.value)
                   for ws in libro.worksheets for fila in ws.iter_rows() for c in fila)


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    p.add_argument("--papel", help="un papel ya generado, para comprobar su apartado de uso de IA")
    args = p.parse_args()

    fallos = 0
    with tempfile.TemporaryDirectory(prefix="probar_cuestionario_") as tmp:
        for nombre, prueba in casos(Path(tmp)):
            try:
                bien = bool(prueba())
            except Exception as exc:                       # noqa: BLE001
                bien, nombre = False, nombre + "  [" + type(exc).__name__ + ": " + str(exc)[:80] + "]"
            print(("  ok   " if bien else "  FALLA ") + nombre)
            fallos += 0 if bien else 1

    if args.papel:
        print("\nPapel: " + args.papel)
        for nombre, prueba in (
                ("(a) lleva el apartado de uso de IA", lambda: tiene_apartado(args.papel)),
                ("(b) huecos de la aplicacion: " + ", ".join(HUECOS),
                 lambda: huecos_docx(args.papel) == HUECOS),
                ("(c) los 11 apartados del equipo vacios", lambda: celdas_equipo(args.papel) == (11, 0))):
            try:
                bien = bool(prueba())
            except Exception as exc:                       # noqa: BLE001
                bien, nombre = False, nombre + "  [" + type(exc).__name__ + ": " + str(exc)[:80] + "]"
            print(("  ok   " if bien else "  FALLA ") + nombre)
            fallos += 0 if bien else 1

    print("\n" + ("Todo bien." if not fallos else str(fallos) + " comprobación(es) fallan."))
    return 0 if not fallos else 1


if __name__ == "__main__":
    sys.exit(main())
