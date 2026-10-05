"""Escribe el papel de trabajo de la prueba MUM en Excel.

    python generar_papel.py --muestra muestra.json --parametros parametros.json \
        --facturas facturas.json [--evaluacion evaluacion.json] \
        --salida "<expediente>/AsistenteIA/FspMum/MUM <PRUEBA> <CLIENTE> <EJERCICIO>.xlsx" \
        --generado 2026-09-03

Una sola hoja, «Analisis muestra»: una fila por elemento seleccionado, con las
columnas repartidas en cuatro ZONAS de color, cada una con su banda de titulo:

  A · Datos de la muestra   lo que ForSampling tiene guardado del elemento
  B · Datos del documento   lo leido en la factura: fecha, CIF, tercero,
                            concepto, base, IVA, IRPF y total
  C · Prueba de muestreo    el campo principal repetido, el valor segun
                            auditoria, el error, el % de error, con que se
                            compara y la observacion
  D · Evidencia del cruce   por que se atribuyo ese documento a ese elemento

El campo principal se repite en la zona C a proposito: la comparacion se lee de
izquierda a derecha aunque la poblacion traiga veinte columnas. Se repite como
REFERENCIA a la celda de la zona A, no como copia, para que el saldo siga
teniendo un solo origen.

El error y el % de error SI van como formula de Excel -lo demas, como valor-.
El motivo es del auditor: cuando cambia el valor segun auditoria, el error y su
porcentaje se recalculan solos, que es justo lo que va a hacer al revisar el
papel. La misma cifra se calcula ademas en Python para lo que se imprime al
ejecutar y para comparar con lo que el auditor ya tenga evaluado.

  Error = Saldo - SaldoAuditoria, la definicion de ForSampling, medida en el
  esquema y respetada tal cual. Positivo = de mas en libros. El % es
  |Error| / |Saldo| x 100, tambien como lo guarda ForSampling, asi que siempre
  sale positivo.

Las dos van envueltas en IFERROR: si el auditor deja la celda con texto, o el
campo principal se queda a cero, la celda se ve EN BLANCO en vez de con un
#!DIV/0! o un #!VALOR! en un papel de trabajo. Y se escribe IFERROR, en ingles,
porque es como el fichero guarda las funciones; Excel la ensena como SI.ERROR
segun el idioma. Escribir SI.ERROR aqui rompe la formula.

La diferencia de dias tambien es formula, la resta de las dos fechas. Para eso
las fechas se escriben como FECHA y no como texto, que de paso deja ordenar y
filtrar por fecha. El sentido es (fecha en libros - fecha del documento), que es
el que calcula el cruce: negativo = el documento es posterior al apunte.

Un elemento sin medir -sin documento, o con el importe ilegible- se queda con
las dos celdas VACIAS y sin formula. Una formula ahi restaria de una celda vacia
y pintaria el saldo entero como error.

La celda del fichero es un HIPERVINCULO al documento, con RUTA ABSOLUTA: el
papel vive en AsistenteIA y los escaneos en otra carpeta, asi que una ruta
relativa se rompe en cuanto el fichero se mueve o se abre desde otro sitio.

La ruta sale del manifiesto de preparar_documentos.py (--manifiesto), que ya
guarda la de cada documento. Y hay un --carpeta-documentos para el caso de
Cowork: alli los scripts corren en un contenedor y la ruta que ven no existe en
la maquina del auditor, asi que el vinculo se rehace sobre la carpeta de Windows
que indique el usuario. Sin ninguna de las dos, la celda queda como texto: mas
vale sin vinculo que con un vinculo que miente.

Un solo color con significado: amarillo en la observacion cuando hay error o
cuando el elemento no se ha podido medir. Los tonos de zona son estructura, y el
azul subrayado del fichero es la convencion de enlace de toda la vida.

**No hay fila de totales, y es a proposito.** Sumar los errores de una MUM es
proyectar a ojo, y sumarlos con su signo esconde las incorrecciones al
cancelarse. La proyeccion la hace ForSampling con la muestra evaluada. Lo que se
imprime al ejecutar son los errores por exceso y por defecto POR SEPARADO.

Y no lee el reloj: la fecha de generacion entra por --generado, para que el
papel se pueda regenerar identico dentro de dos años.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib_fsp import (  # noqa: E402
    rutas_documentos,
    TOL_IMPORTE,
    cargar_evaluacion,
    cargar_json,
    cargar_facturas,
    cargar_muestra,
    cargar_parametros,
    cruzar,
    columnas_de_muestra,
    parse_fecha,
    parse_importe,
    salida_utf8,
)
from lib_mum import (  # noqa: E402
    NOMBRE_TERMINO,
    comparar_con_auditor_mum,
    evaluar_mum,
    observacion_mum,
    recuento,
    termino_mayoritario,
)

FONT = "Arial"
FORMATO_EURO = '#,##0.00 "€";-#,##0.00 "€";"-"'
FORMATO_PCT = '0.00 "%";-0.00 "%";"-"'
FORMATO_FECHA = "DD/MM/YYYY"
AMARILLO = PatternFill("solid", fgColor="FFFF00")
AZUL_ENLACE = "0563C1"

# Una zona por bloque de columnas: (titulo de la banda, color de cabecera, tinte
# de los datos). El tinte es flojo a proposito, para que el amarillo de la
# observacion -que es el unico color con significado- siga cantando encima.
ZONAS_DEF = {
    "muestra": ("A · Datos de la muestra", "1F4E78", None),
    "documento": ("B · Datos del documento", "217346", "EAF3EC"),
    "prueba": ("C · Prueba de muestreo", "C55A11", "FDF0E6"),
    "cruce": ("D · Evidencia del cruce", "7F7F7F", "F2F2F2"),
}

# «Proveedor o cliente» lleva lo que el documento diga: la razon social si la factura va en
# claro, y el TOKEN del sello si va tachada -que `rehidratar` convierte en el nombre real al
# entregar-. El CIF, con la factura tachada, se queda vacio: esta en negro y no se restaura,
# porque el diccionario indexa por cuenta y no por CIF.
COL_DOCUMENTO = ["Fichero", "Nº factura", "Fecha doc.", "CIF", "Proveedor o cliente", "Concepto",
                 "Base", "IVA", "IRPF", "Total"]
COL_CRUCE = ["Casa por", "Días libros–doc."]

_BORDE_LADO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_BORDE_LADO, right=_BORDE_LADO, top=_BORDE_LADO, bottom=_BORDE_LADO)


def _bandas(ws, fila: int, tramos: list[tuple[str, str, int, int]]) -> None:
    """La fila de titulos de zona, cada uno combinado sobre sus columnas."""
    for titulo, color, desde, hasta in tramos:
        for j in range(desde, hasta + 1):
            c = ws.cell(row=fila, column=j)
            c.fill = PatternFill("solid", fgColor=color)
            c.border = BORDE
        if hasta > desde:
            ws.merge_cells(start_row=fila, start_column=desde, end_row=fila, end_column=hasta)
        c = ws.cell(row=fila, column=desde, value=titulo)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center")


def _cabecera(ws, fila: int, textos: list[str], colores: list[str]) -> None:
    for j, (t, color) in enumerate(zip(textos, colores), start=1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=color)
        c.border = BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _anchos(ws, anchos: dict[int, float]) -> None:
    for col, w in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _eur(v: float | None) -> str:
    if v is None:
        return "—"
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _hoja(wb: Workbook, evaluados: list[dict], cols: dict, params: dict, generado: str,
          rutas: dict[str, str] | None = None) -> None:
    """La unica hoja: un elemento por fila, en cuatro zonas de color.

    Las tres columnas de la MUM SI se rellenan, al contrario que los atributos de
    la prueba de cumplimiento, y la diferencia no es un descuido. Un atributo es
    un veredicto -Ok o no Ok- y rellenarlo es concluir, que no es del skill. Un
    importe segun el documento es una medida: decir que la factura pone 16.962,50
    es describir. De quien es la propuesta lo dice la linea de arriba de la hoja
    -«Propuesta del asistente... el auditor adopta o cambia cada importe»- y el
    propio rotulo «Valor Auditoria (doc)», que dice de donde sale el importe.
    """
    ws = wb.active
    ws.title = "Análisis muestra"
    if not evaluados:
        ws["A1"] = "La muestra no trae elementos."
        return

    pr = params.get("parametros") or {}
    et = parse_importe(pr.get("ErrorTolerableValor"))

    def _dia(v):
        d = parse_fecha(v)
        return d.strftime("%d/%m/%Y") if d else (str(v) if v else "—")

    col_pob = [c for c in evaluados[0]["fila"].keys() if c.lower() not in ("seleccionado",)]
    col_principal = cols.get("importe")
    # «VRL» es el valor registrado en libros, el termino de la MUM. Se rotula asi y no
    # con el nombre de la columna de ForSampling -que cambia de prueba en prueba- porque
    # el auditor lee la zona C en el lenguaje del muestreo, no en el del fichero.
    col_mum = ["VRL (muestra)", "Valor Auditoría (doc)", "Error",
               "% error", "Comparado con", "Observación propuesta"]
    cab = col_pob + COL_DOCUMENTO + col_mum + COL_CRUCE

    # Donde empieza cada zona, para las bandas y para las formulas
    n_pob, n_doc, n_mum = len(col_pob), len(COL_DOCUMENTO), len(col_mum)
    ini_doc = n_pob + 1
    ini_mum = n_pob + n_doc + 1
    ini_cruce = ini_mum + n_mum
    c_saldo, c_valor, c_error, c_pct = ini_mum, ini_mum + 1, ini_mum + 2, ini_mum + 3
    c_obs = ini_mum + n_mum - 1
    # el campo principal de la zona A, para referenciarlo desde la C
    col_saldo_pob = (col_pob.index(col_principal) + 1) if col_principal in col_pob else None
    col_fecha_pob = ((col_pob.index(cols['fecha']) + 1)
                     if cols.get('fecha') in col_pob else None)
    c_fecha_doc, c_dias = ini_doc + 2, len(cab)

    colores = ([ZONAS_DEF["muestra"][1]] * n_pob + [ZONAS_DEF["documento"][1]] * n_doc
               + [ZONAS_DEF["prueba"][1]] * n_mum + [ZONAS_DEF["cruce"][1]] * len(COL_CRUCE))
    tintes = ([ZONAS_DEF["muestra"][2]] * n_pob + [ZONAS_DEF["documento"][2]] * n_doc
              + [ZONAS_DEF["prueba"][2]] * n_mum + [ZONAS_DEF["cruce"][2]] * len(COL_CRUCE))

    ws["A1"] = f"Análisis de la muestra — {params.get('Prueba') or 'prueba MUM'}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = (f"Prueba {params.get('MuestraId')} · MUM · área {params.get('Area')} · "
                f"ref. {params.get('Referencia')} · ejercicio "
                f"{_dia(params.get('FechaInicioAuditoria'))} — {_dia(params.get('FechaFinAuditoria'))} · "
                f"unidad de muestreo {pr.get('UnidadMuestreo') or '—'} · población "
                f"{pr.get('PoblacionNumElementos') or '—'} elementos · error tolerable de la prueba "
                f"{_eur(et)} · generado {generado}")
    ws["A2"].font = Font(name=FONT, size=9)
    ws["A3"] = ("Propuesta del asistente: localiza el documento de cada elemento y mide lo que sostiene, "
                "en el término con el que contabiliza esta población. El auditor adopta o cambia cada importe. "
                "El error y el % son fórmulas: al cambiar el valor según auditoría se recalculan.")
    ws["A3"].font = Font(name=FONT, italic=True, size=9, color="595959")
    ws["A4"] = ("Error = campo principal − valor según auditoría, como lo define ForSampling: positivo es de "
                "más en libros. Amarillo: hay diferencia, o el elemento no se ha podido medir. Este papel NO "
                "proyecta el error a la población, no lo compara con el error tolerable y no suma los errores "
                "entre sí: eso lo hace ForSampling, y los errores por exceso y por defecto no se cancelan.")
    ws["A4"].font = Font(name=FONT, italic=True, size=9, color="595959")

    _bandas(ws, 6, [(ZONAS_DEF["muestra"][0], ZONAS_DEF["muestra"][1], 1, n_pob),
                    (ZONAS_DEF["documento"][0], ZONAS_DEF["documento"][1], ini_doc, ini_doc + n_doc - 1),
                    (ZONAS_DEF["prueba"][0], ZONAS_DEF["prueba"][1], ini_mum, ini_mum + n_mum - 1),
                    (ZONAS_DEF["cruce"][0], ZONAS_DEF["cruce"][1], ini_cruce, len(cab))])
    _cabecera(ws, 7, cab, colores)

    L = get_column_letter
    fila = 8
    for e in evaluados:
        f, fac, c = e["fila"], e["factura"], e["criterios"]
        valores = []
        # Las fechas van como FECHA, no como texto: es lo que permite que la
        # diferencia de dias sea una resta de celdas, y de paso que se puedan
        # ordenar y filtrar por fecha. Si alguna no se puede interpretar se deja
        # el texto tal cual y esa fila se queda sin formula.
        fecha_libros = fecha_doc = None
        for k in col_pob:
            v = f.get(k)
            if k == col_principal:
                v = parse_importe(v)
            elif k == cols.get("fecha"):
                fecha_libros = parse_fecha(v)
                v = fecha_libros or v
            valores.append(v)

        if fac:
            fecha_doc = parse_fecha(fac.get("fecha"))
            valores += [fac.get("fichero"), fac.get("numero"),
                        fecha_doc or fac.get("fecha"), fac.get("cif"),
                        fac.get("proveedor") or fac.get("token"), fac.get("concepto"),
                        parse_importe(fac.get("base")), parse_importe(fac.get("iva")),
                        parse_importe(fac.get("irpf")), parse_importe(fac.get("total"))]
        else:
            valores += ["— no localizada", None, None, None, None, None, None, None, None, None]

        # Zona C. El campo principal se referencia, no se copia.
        valores.append(f"={L(col_saldo_pob)}{fila}" if col_saldo_pob else None)
        valores.append(e["saldo_auditoria"])
        if e["saldo_auditoria"] is not None and col_saldo_pob:
            valores.append(f'=IFERROR({L(c_saldo)}{fila}-{L(c_valor)}{fila},"")')
            valores.append(f'=IFERROR(ABS({L(c_error)}{fila})'
                           f'/ABS({L(c_saldo)}{fila})*100,"")')
        else:
            # sin medir: las dos celdas se quedan vacias. Una formula aqui restaria
            # de una celda vacia y pintaria el saldo entero como error.
            valores += [None, None]
        valores.append(NOMBRE_TERMINO.get(e["termino"], "—"))
        valores.append(e["observacion"])

        if fac:
            casa = ", ".join(x for x, b in (("número", c["numero"]), ("tercero", c.get("tercero")),
                                            (f"importe ({c['importe']})", c["importe"]),
                                            ("fecha", c["fecha"])) if b)
            if fecha_libros and fecha_doc:
                dias = f'=IFERROR({L(col_fecha_pob)}{fila}-{L(c_fecha_doc)}{fila},"")'
            else:
                dias = c["dias"]   # sin dos fechas de verdad no hay resta que hacer
            valores += [casa, dias]
        else:
            valores += ["", None]

        for j, v in enumerate(valores, start=1):
            cell = ws.cell(row=fila, column=j, value=v)
            cell.font = Font(name=FONT, size=9)
            cell.border = BORDE
            cell.alignment = Alignment(vertical="top", wrap_text=j == c_obs)
            if tintes[j - 1]:
                cell.fill = PatternFill("solid", fgColor=tintes[j - 1])
            if isinstance(v, float):
                cell.number_format = FORMATO_EURO
        # el fichero, enlazado al documento. Despues del bucle, que fija la fuente.
        enlace = (rutas or {}).get((fac or {}).get("fichero"))
        if enlace:
            cel = ws.cell(row=fila, column=ini_doc)
            cel.hyperlink = enlace
            cel.font = Font(name=FONT, size=9, color=AZUL_ENLACE, underline="single")
        for col in (c_saldo, c_valor, c_error):
            ws.cell(row=fila, column=col).number_format = FORMATO_EURO
        ws.cell(row=fila, column=c_pct).number_format = FORMATO_PCT
        ws.cell(row=fila, column=c_dias).number_format = "0"
        for col in (col_fecha_pob, c_fecha_doc):
            if col:
                ws.cell(row=fila, column=col).number_format = FORMATO_FECHA
        if e["error"] is None or abs(e["error"]) > TOL_IMPORTE:
            ws.cell(row=fila, column=c_obs).fill = AMARILLO
        fila += 1

    ws.auto_filter.ref = f"A7:{L(len(cab))}{fila - 1}"
    ws.freeze_panes = "A8"
    anchos = {i + 1: 14 for i in range(len(cab))}
    for i, k in enumerate(col_pob, start=1):
        if k in (cols.get("tercero"), "Acreedor"):
            anchos[i] = 30
    anchos[ini_doc] = 34                      # Fichero
    anchos[ini_doc + 4] = 30                  # Proveedor o cliente
    anchos[ini_doc + 5] = 40                  # Concepto
    for k in range(n_mum - 1):
        anchos[ini_mum + k] = 20
    anchos[c_obs] = 95
    anchos[ini_cruce] = 26                    # Casa por
    _anchos(ws, anchos)


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--muestra", required=True)
    p.add_argument("--parametros", required=True)
    p.add_argument("--facturas", required=True)
    p.add_argument("--evaluacion", help="exportar_consulta(entidad='evaluacion', id=N), para comparar")
    p.add_argument("--manifiesto", help="manifiesto.json de preparar_documentos.py: de ahí sale la "
                                        "ruta de cada documento para el hipervínculo")
    p.add_argument("--carpeta-documentos", dest="carpeta_documentos",
                   help="carpeta de los escaneos EN LA MÁQUINA DEL AUDITOR. Manda sobre el "
                        "manifiesto: en Cowork la ruta del contenedor no le sirve de nada")
    p.add_argument("--salida", required=True)
    p.add_argument("--generado", required=True, help="fecha de generación, AAAA-MM-DD. Nada lee el reloj")
    args = p.parse_args()

    muestra = cargar_muestra(args.muestra)
    params = cargar_parametros(args.parametros)
    facturas = cargar_facturas(args.facturas)
    evaluacion = cargar_evaluacion(args.evaluacion)

    cols, _, _ = columnas_de_muestra(muestra)
    cruce = cruzar(muestra, facturas, cols)
    evaluados = evaluar_mum(cruce, cols)
    for e in evaluados:
        e["observacion"] = observacion_mum(e)
    comp = comparar_con_auditor_mum(evaluados, (evaluacion or {}).get("filas") or [], cols)
    r = recuento(evaluados)

    rutas = rutas_documentos(args.manifiesto, args.carpeta_documentos)

    wb = Workbook()
    _hoja(wb, evaluados, cols, params, args.generado, rutas)

    salida = Path(args.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(salida)
    except PermissionError:
        print(f"ERROR: no se puede escribir {salida}. Si está abierto en Excel, ciérralo y repite.")
        return 2

    term = termino_mayoritario(evaluados)
    print(f"Papel escrito: {salida}")
    if rutas:
        enlazados = [e for e in evaluados if e["factura"] and rutas.get(e["factura"].get("fichero"))]
        rotos = [e for e in enlazados if not Path(rutas[e["factura"]["fichero"]]).exists()]
        print(f"  documentos enlazados {len(enlazados)} de {r['con_documento']}")
        if rotos:
            print(f"  AVISO: {len(rotos)} vínculo(s) apuntan a rutas que este proceso no ve. Si los "
                  f"scripts corren en un contenedor es lo normal y en la máquina del auditor "
                  f"funcionarán; si no, pásale --carpeta-documentos con la ruta buena y repite")
    else:
        print("  sin hipervínculos: no se ha indicado --manifiesto ni --carpeta-documentos")
    print(f"  elementos {r['elementos']} · unidades de muestreo con repeticiones {r['repeticiones']:g}"
          f" · con documento {r['con_documento']} · sin documento {r['sin_documento']}"
          f" · medidos {r['medidos']} · sin medir {r['sin_medir']}")
    print(f"  esta población contabiliza por: {NOMBRE_TERMINO.get(term, 'sin criterio claro (lo decide el auditor)')}")
    print(f"  sin diferencia {r['sin_error']} · con diferencia {r['con_error']}")
    if r["con_error"]:
        print(f"  incorrecciones POR SEPARADO, sin netear: {r['n_exceso']} por exceso suman "
              f"{_eur(r['suma_exceso'])} · {r['n_defecto']} por defecto suman {_eur(r['suma_defecto'])}")
        print("  NO son la proyección ni el error neto: la proyección la hace ForSampling con la muestra evaluada")
    for e in evaluados:
        if e["error"] is None or abs(e["error"]) > TOL_IMPORTE:
            ident = e["fila"].get(cols.get("id"), "?")
            print(f"  · elemento {ident}: {e['observacion']}")
    for par in cruce.get("candidatos_sueltos") or []:
        print(f"  · POSIBLE DIFERENCIA, no extravío: el elemento {par['id']} quedó sin documento y "
              f"{par['fichero']} sin elemento"
              + (", mismo tercero" if par["tercero"] else "")
              + (f", {par['dias']} día(s) de diferencia" if par["dias"] is not None else "")
              + (f". En libros {par['importe_libros']}, en el documento base {par['base']} / total "
                 f"{par['total']}: diferencia mínima {par['diferencia_minima']}"
                 if par["diferencia_minima"] is not None else "")
              + ". Revísalo a mano: si son el mismo, hay una diferencia real que declarar con "
                "poblacion_id en facturas.json")
    for fac in cruce["facturas_sin_fila"]:
        print(f"  · documento que no es de ningún elemento de la muestra: {fac.get('fichero')}"
              f" ({fac.get('proveedor') or '—'}, {fac.get('total') or '—'})")
    if comp:
        c = comp["recuento"]
        print(f"  frente al auditor: coinciden {c['coinciden']} · skill señala error/auditor 0 "
              f"{c['skill_error_auditor_cero']} · SKILL 0/AUDITOR ERROR {c['skill_cero_auditor_error']}"
              f" · los dos con error distinto {c['discrepan']} · sin comparar {c['sin_evaluar']}")
        for f in comp["filas"]:
            if f["estado"].startswith("EL SKILL DA 0") or f["estado"].startswith("los dos"):
                print(f"    · elemento {f['id']}: skill {_eur(f['error_skill'])} / auditor "
                      f"{_eur(f['error_auditor'])} — {f['observacion_auditor'] or ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
